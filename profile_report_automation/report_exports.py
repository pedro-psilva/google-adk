from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from profile_report_automation.neopi_language import NEOPI_DOMAIN_ORDER, rewrite_neopi_synthesis_text
from profile_report_automation.pdf_export import (
    PDF_EXPORT_WORKSHEET_NAME,
    export_workbook_sheet_via_excel,
    export_workbook_sheet_via_worker,
    get_pdf_export_worker_url,
)


METHODOLOGY_TEXT = (
    "Este documento apresenta uma sintese do perfil comportamental em contexto organizacional, "
    "a partir da leitura integrada dos instrumentos aplicados. Os resultados devem ser considerados "
    "como indicadores para reflexao e desenvolvimento, sempre observados em conjunto com o contexto "
    "da pessoa, da funcao e do momento profissional."
)

SECTION_INTROS = {
    "neopi": [
        "O Inventario de Personalidade NEO PI-R e uma ferramenta baseada no modelo dos Cinco Grandes Fatores, "
        "utilizada para apoiar a leitura de tendencias comportamentais relevantes no ambiente de trabalho."
    ],
    "profiler": [
        "O Profiler apoia o mapeamento do estilo comportamental predominante a partir de quatro eixos: "
        "Executor, Comunicador, Planejador e Analista."
    ],
    "career_anchors": [
        "As Ancoras de Carreira ajudam a identificar valores, motivadores e preferencias de percurso profissional."
    ],
    "cultural_diagnosis": [
        "O Diagnostico Cultural indica os ambientes organizacionais com maior aderencia percebida pela pessoa avaliada."
    ],
}


def export_local_reports(
    output_dir: str | Path,
    bundle: dict[str, Any],
    coverage: dict[str, Any],
    draft: dict[str, Any],
    *,
    workbook_path: str | Path | None = None,
) -> dict[str, str]:
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)

    base_name = build_report_basename(bundle)
    docx_path = target_dir / f"{base_name}.docx"
    pdf_path = target_dir / f"{base_name}.pdf"

    build_docx_report(docx_path, bundle, draft)
    if not workbook_path:
        raise RuntimeError(
            "A geracao de PDF exige a planilha final preenchida. O PDF alternativo foi desativado."
        )
    build_pdf_from_workbook(pdf_path, workbook_path)

    return {
        "docx": str(docx_path.resolve()),
        "pdf": str(pdf_path.resolve()),
    }


def build_report_basename(bundle: dict[str, Any]) -> str:
    raw_name = str(bundle.get("person", {}).get("name") or "Pessoa avaliada")
    normalized = unicodedata.normalize("NFKD", raw_name)
    without_marks = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    safe_name = re.sub(r'[<>:"/\\\\|?*]+', " ", without_marks)
    safe_name = re.sub(r"\s+", " ", safe_name).strip() or "Pessoa avaliada"
    return f"{safe_name} - Relatorio de Analise de Perfil"


def build_docx_report(path: Path, bundle: dict[str, Any], draft: dict[str, Any]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)

    _add_docx_title(document, "ANALISE DE PERFIL COMPORTAMENTAL")
    _add_docx_heading(document, "I. Identificacao:")
    _add_docx_identification(document, bundle)

    _add_docx_heading(document, "II. Metodologia e objetivo:")
    document.add_paragraph(METHODOLOGY_TEXT)

    _add_docx_heading(document, "III. Resultados da Analise:")
    draft_sections = _draft_sections_by_key(draft)

    _add_docx_subsection(document, "3.1- NEO PI-R", SECTION_INTROS["neopi"], _build_neopi_paragraphs(bundle, draft_sections))
    _add_docx_simple_table(
        document,
        ["Fator", "Categoria", "T score"],
        [
            [item["domain"], _title_case(str(item["category"])), str(item["t_score"])]
            for item in bundle.get("neopi", {}).get("domains", [])
            if item.get("citation_candidate")
        ],
    )

    _add_docx_subsection(
        document,
        "3.2- Profiler",
        SECTION_INTROS["profiler"] + _build_profiler_intro(bundle),
        _build_section_paragraphs("profiler", bundle, draft_sections),
    )
    _add_docx_simple_table(
        document,
        ["Estilo", "Percentual"],
        [
            [item["style"], f"{item['percentage']:.2f}%"]
            for item in bundle.get("profiler", {}).get("scores", [])
        ],
    )

    _add_docx_subsection(
        document,
        "3.3- Ancoras de Carreira",
        SECTION_INTROS["career_anchors"],
        _build_section_paragraphs("career_anchors", bundle, draft_sections),
    )
    _add_docx_simple_table(
        document,
        ["Ancora", "Media"],
        [
            [item["name"], str(item["average"])]
            for item in bundle.get("career_anchors", {}).get("scores", [])
        ],
    )

    _add_docx_subsection(
        document,
        "3.4- Diagnostico Cultural",
        SECTION_INTROS["cultural_diagnosis"],
        _build_section_paragraphs("cultural_diagnosis", bundle, draft_sections),
    )
    _add_docx_simple_table(
        document,
        ["Cultura", "Pontuacao"],
        [
            [item["culture"], f"{item['score']:.2f}"]
            for item in bundle.get("cultural_diagnosis", {}).get("scores", [])
        ],
    )

    _add_docx_heading(document, "IV. Conclusao:")
    document.add_paragraph("A partir dos indicadores de seu perfil, destacam-se os seguintes pontos:")
    _add_docx_bullets(document, _build_conclusion_bullets(bundle, draft_sections))

    document.save(path)


def build_pdf_report(path: Path, bundle: dict[str, Any], draft: dict[str, Any]) -> None:
    raise RuntimeError(
        "A geracao de PDF alternativo foi desativada. O PDF deve ser exportado a partir da planilha final."
    )


def build_pdf_from_workbook(path: Path, workbook_path: str | Path) -> None:
    if export_workbook_sheet_via_excel(path, workbook_path, worksheet_name=PDF_EXPORT_WORKSHEET_NAME):
        return
    if export_workbook_sheet_via_worker(path, workbook_path, worksheet_name=PDF_EXPORT_WORKSHEET_NAME):
        return

    worker_url = get_pdf_export_worker_url()
    if worker_url:
        raise RuntimeError(
            "Nao foi possivel exportar o PDF com fidelidade ao Excel, mesmo com o worker configurado. "
            "Verifique se o worker Windows esta ativo e se o Microsoft Excel esta disponivel para automacao."
        )

    raise RuntimeError(
        "Nao foi possivel exportar o PDF com fidelidade ao Excel. Execute o backend em Windows com Microsoft Excel "
        "ou inicie o worker Windows configurando PDF_EXPORT_WORKER_URL."
    )


def _draw_pdf_footer(canvas: Any, document: Any) -> None:
    canvas.saveState()
    canvas.setFont("ReportBody" if "ReportBody" in pdfmetrics.getRegisteredFontNames() else "Helvetica", 9)
    canvas.setFillColor(colors.HexColor("#6b7280"))
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Pagina {canvas.getPageNumber()}")
    canvas.restoreState()


def _build_pdf_styles() -> dict[str, ParagraphStyle]:
    sample = getSampleStyleSheet()
    registered_fonts = set(pdfmetrics.getRegisteredFontNames())
    body_font = "ReportBody" if "ReportBody" in registered_fonts else "Helvetica"
    heading_font = "ReportHeadingBold" if "ReportHeadingBold" in registered_fonts else "Helvetica-Bold"
    return {
        "title": ParagraphStyle(
            "ReportTitle",
            parent=sample["Title"],
            fontName=heading_font,
            fontSize=16,
            leading=20,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#202124"),
            spaceAfter=6,
        ),
        "heading": ParagraphStyle(
            "ReportHeading",
            parent=sample["Heading2"],
            fontName=heading_font,
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#111827"),
            spaceBefore=2,
            spaceAfter=6,
        ),
        "subheading": ParagraphStyle(
            "ReportSubheading",
            parent=sample["Heading3"],
            fontName=heading_font,
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#111827"),
            spaceAfter=4,
        ),
        "sheetheading": ParagraphStyle(
            "ReportSheetHeading",
            parent=sample["Heading2"],
            fontName=heading_font,
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=6,
        ),
        "body": ParagraphStyle(
            "ReportBodyStyle",
            parent=sample["BodyText"],
            fontName=body_font,
            fontSize=10.2,
            leading=14,
            textColor=colors.HexColor("#1f2937"),
            spaceAfter=5,
        ),
        "label": ParagraphStyle(
            "ReportLabel",
            parent=sample["BodyText"],
            fontName=heading_font,
            fontSize=10.2,
            leading=13,
            textColor=colors.HexColor("#111827"),
            spaceAfter=2,
        ),
        "footer": ParagraphStyle(
            "ReportFooter",
            parent=sample["BodyText"],
            fontName=body_font,
            fontSize=9.6,
            leading=12,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#4b5563"),
            spaceAfter=2,
        ),
        "sheetbody": ParagraphStyle(
            "ReportSheetBody",
            parent=sample["BodyText"],
            fontName=body_font,
            fontSize=9.1,
            leading=11.2,
            textColor=colors.HexColor("#1f2937"),
            spaceAfter=1.2,
        ),
        "sheetsubheading": ParagraphStyle(
            "ReportSheetSubheading",
            parent=sample["Heading3"],
            fontName=heading_font,
            fontSize=10,
            leading=11.8,
            textColor=colors.HexColor("#111827"),
            spaceAfter=1.4,
        ),
    }


def _register_pdf_fonts() -> None:
    windows_fonts = Path("C:/Windows/Fonts")
    regular_font = windows_fonts / "arial.ttf"
    bold_font = windows_fonts / "arialbd.ttf"

    if regular_font.exists() and "ReportBody" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ReportBody", str(regular_font)))
    if bold_font.exists() and "ReportHeadingBold" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("ReportHeadingBold", str(bold_font)))


def _add_docx_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(15.5)
    run.font.name = "Arial"


def _add_docx_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.name = "Arial"


def _add_docx_identification(document: Document, bundle: dict[str, Any]) -> None:
    person = bundle.get("person", {})
    rows = [
        ("Nome", str(person.get("name") or "-")),
        ("Data da aplicacao", str(person.get("application_date") or "-")),
        ("Unidade de negocios", str(person.get("business_unit") or "-")),
        ("Demanda", str(person.get("demand") or "-")),
        ("Cargo", str(person.get("role") or "-")),
        ("Procedimentos realizados", _build_procedures_label(bundle)),
    ]

    for label, value in rows:
        paragraph = document.add_paragraph()
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        paragraph.add_run(value)


def _add_docx_subsection(document: Document, title: str, intro_lines: list[str], paragraphs: list[str]) -> None:
    heading = document.add_paragraph()
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = "Arial"

    for line in intro_lines:
        if line:
            document.add_paragraph(line)
    for paragraph in paragraphs:
        if paragraph:
            document.add_paragraph(paragraph)


def _add_docx_simple_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    if not rows:
        return

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value

    document.add_paragraph()


def _add_docx_bullets(document: Document, bullets: list[str]) -> None:
    for item in bullets:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def _build_pdf_identification(bundle: dict[str, Any], styles: dict[str, ParagraphStyle]) -> list[Any]:
    person = bundle.get("person", {})
    rows = [
        f"<b>Nome:</b> {escape(str(person.get('name') or '-'))}",
        f"<b>Data da aplicacao:</b> {escape(str(person.get('application_date') or '-'))}",
        f"<b>Unidade de negocios:</b> {escape(str(person.get('business_unit') or '-'))}",
        f"<b>Demanda:</b> {escape(str(person.get('demand') or '-'))}",
        f"<b>Cargo:</b> {escape(str(person.get('role') or '-'))}",
        f"<b>Procedimentos realizados:</b> {escape(_build_procedures_label(bundle))}",
    ]
    return [Paragraph(row, styles["body"]) for row in rows]


def _build_pdf_subsection(
    title: str,
    intro_lines: list[str],
    paragraphs: list[str],
    styles: dict[str, ParagraphStyle],
) -> list[Any]:
    items: list[Any] = [Paragraph(title, styles["subheading"])]
    items.extend(Paragraph(escape(line), styles["body"]) for line in intro_lines if line)
    items.extend(Paragraph(escape(line), styles["body"]) for line in paragraphs if line)
    return items


def _build_pdf_table(headers: list[str], rows: list[list[str]]) -> Table:
    data = [headers, *rows] if rows else [headers]
    body_font = "ReportBody" if "ReportBody" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
    heading_font = "ReportHeadingBold" if "ReportHeadingBold" in pdfmetrics.getRegisteredFontNames() else "Helvetica-Bold"
    table = Table(data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111827")),
                ("FONTNAME", (0, 0), (-1, 0), heading_font),
                ("FONTNAME", (0, 1), (-1, -1), body_font),
                ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _extract_nonempty_worksheet_rows(worksheet: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in worksheet.iter_rows():
        values = [_format_worksheet_cell(cell) for cell in row]
        compact_values = [value for value in values if value]
        if compact_values:
            rows.append(compact_values)
    return rows


def _format_worksheet_cell(cell: Any) -> str:
    value = cell.value
    if value is None:
        return ""

    if hasattr(value, "strftime"):
        try:
            return value.strftime("%d/%m/%Y")
        except TypeError:
            return str(value).strip()

    if isinstance(value, bool):
        return "Sim" if value else "Nao"

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number_format = str(cell.number_format or "")
        if "%" in number_format:
            decimals = _percentage_decimal_places(number_format)
            return f"{value * 100:.{decimals}f}%"
        if isinstance(value, float):
            return f"{value:.2f}".rstrip("0").rstrip(".")
        return str(value)

    return str(value).strip()


def _percentage_decimal_places(number_format: str) -> int:
    match = re.search(r"[.,](0+)%", number_format)
    if not match:
        return 0
    return len(match.group(1))


def _build_pdf_story_from_worksheet_rows(
    rows: list[list[str]],
    styles: dict[str, ParagraphStyle],
    *,
    honor_page_break_markers: bool = False,
) -> list[Any]:
    story: list[Any] = []
    pending_table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal pending_table_rows
        if not pending_table_rows:
            return
        story.append(_build_workbook_pdf_table(pending_table_rows, styles))
        story.append(Spacer(1, 3 * mm))
        pending_table_rows = []

    total_rows = len(rows)
    for row_index, row in enumerate(rows):
        if len(row) == 1:
            flush_table()
            text = row[0]
            if honor_page_break_markers and _is_page_break_marker(text):
                footer_lines = [line.strip() for line in text.splitlines() if line.strip()]
                if len(footer_lines) > 1:
                    story.append(Spacer(1, 2 * mm))
                    for footer_line in footer_lines:
                        story.append(Paragraph(_paragraph_text(footer_line), styles["footer"]))
                if row_index < total_rows - 1:
                    story.append(PageBreak())
                continue
            paragraph_style = styles["sheetsubheading"] if _looks_like_heading(text) else styles["sheetbody"]
            story.append(Paragraph(_paragraph_text(text), paragraph_style))
            if _looks_like_heading(text):
                story.append(Spacer(1, 0.5 * mm))
            continue

        pending_table_rows.append(row)

    flush_table()
    return story


def _build_workbook_pdf_table(
    rows: list[list[str]],
    styles: dict[str, ParagraphStyle],
) -> Table:
    max_columns = max(len(row) for row in rows)
    padded_rows = [row + [""] * (max_columns - len(row)) for row in rows]
    header_row = _looks_like_table_header(padded_rows)
    body_font = "ReportBody" if "ReportBody" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
    heading_font = "ReportHeadingBold" if "ReportHeadingBold" in pdfmetrics.getRegisteredFontNames() else "Helvetica-Bold"

    table_data: list[list[Any]] = []
    for row_index, row in enumerate(padded_rows):
        formatted_row: list[Any] = []
        for cell_value in row:
            style = styles["label"] if header_row and row_index == 0 else styles["body"]
            formatted_row.append(Paragraph(_paragraph_text(cell_value), style))
        table_data.append(formatted_row)

    table = Table(
        table_data,
        hAlign="LEFT",
        colWidths=_estimate_workbook_table_widths(padded_rows),
        repeatRows=1 if header_row else 0,
    )
    style_commands = [
        ("FONTNAME", (0, 0), (-1, -1), body_font),
        ("FONTSIZE", (0, 0), (-1, -1), 9.3),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#d1d5db")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header_row:
        style_commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111827")),
                ("FONTNAME", (0, 0), (-1, 0), heading_font),
            ]
        )
    table.setStyle(TableStyle(style_commands))
    return table


def _estimate_workbook_table_widths(rows: list[list[str]]) -> list[float]:
    page_width = A4[0] - (36 * mm)
    column_count = max(len(row) for row in rows)
    if column_count == 0:
        return []

    max_lengths = [0] * column_count
    for row in rows:
        for index, value in enumerate(row):
            max_lengths[index] = max(max_lengths[index], min(len(value), 80))

    weight_floor = 12
    weights = [max(length, weight_floor) for length in max_lengths]
    weight_sum = sum(weights) or column_count
    widths = [(page_width * weight) / weight_sum for weight in weights]

    minimum_width = 26 * mm
    if any(width < minimum_width for width in widths):
        return [page_width / column_count] * column_count
    return widths


def _looks_like_table_header(rows: list[list[str]]) -> bool:
    if len(rows) < 2:
        return False

    first_row = rows[0]
    nonempty_cells = [cell for cell in first_row if cell.strip()]
    if len(nonempty_cells) < 2:
        return False

    average_size = sum(len(cell) for cell in nonempty_cells) / len(nonempty_cells)
    return average_size <= 30


def _looks_like_heading(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    if stripped.endswith(":"):
        return True
    if re.match(r"^[IVXLC]+\.\s", stripped):
        return True
    if re.match(r"^\d+([.-]\d+)*\s*[-:]", stripped):
        return True
    letters = [character for character in stripped if character.isalpha()]
    return bool(letters) and stripped == stripped.upper()


def _paragraph_text(value: str) -> str:
    escaped = escape(value.strip())
    return escaped.replace("\n", "<br/>")


def _preferred_pdf_worksheets(workbook: Any) -> list[Any]:
    visible_sheets = [worksheet for worksheet in workbook.worksheets if worksheet.sheet_state == "visible"]
    summary_sheets = [worksheet for worksheet in visible_sheets if _is_summary_sheet(worksheet.title)]
    return summary_sheets or visible_sheets


def _worksheet_pdf_margins(worksheet: Any) -> tuple[float, float, float, float]:
    margins = worksheet.page_margins
    return (
        float(getattr(margins, "left", 0.7)) * 25.4 * mm,
        float(getattr(margins, "right", 0.7)) * 25.4 * mm,
        float(getattr(margins, "top", 0.75)) * 25.4 * mm,
        float(getattr(margins, "bottom", 0.75)) * 25.4 * mm,
    )


def _is_summary_sheet(title: str) -> bool:
    normalized = unicodedata.normalize("NFKD", title)
    without_marks = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return without_marks.strip().lower() == "sintese"


def _is_page_break_marker(text: str) -> bool:
    first_line = text.splitlines()[0].strip()
    return bool(re.match(r"^P[aá]gina\s+\d+\s+de\s+\d+$", first_line, flags=re.IGNORECASE))


def _draft_sections_by_key(draft: dict[str, Any]) -> dict[str, dict[str, Any]]:
    sections = {
        section["key"]: section
        for section in draft.get("sections", [])
        if isinstance(section, dict) and section.get("key")
    }
    neopi_section = sections.get("neopi")
    if neopi_section is not None and "neopi_factor_summaries" not in neopi_section:
        neopi_section["neopi_factor_summaries"] = draft.get("neopi_factor_summaries", [])
    return sections


def _build_neopi_paragraphs(bundle: dict[str, Any], draft_sections: dict[str, dict[str, Any]]) -> list[str]:
    factor_summaries = _build_neopi_factor_paragraphs(bundle, draft_sections)
    if factor_summaries:
        return factor_summaries

    paragraphs = _build_section_paragraphs("neopi", bundle, draft_sections)
    if paragraphs:
        return paragraphs

    synthesis_map = bundle.get("neopi", {}).get("friendly_synthesis_by_domain") or bundle.get("neopi", {}).get("synthesis_by_domain", {})
    lines: list[str] = []
    for domain_name in NEOPI_DOMAIN_ORDER:
        friendly_text = rewrite_neopi_synthesis_text(str(synthesis_map.get(domain_name) or ""), domain_name=domain_name)
        if friendly_text:
            lines.append(f"{domain_name}: {friendly_text}")
            continue

        for item in bundle.get("neopi", {}).get("domains", []):
            if item.get("domain") != domain_name or not item.get("citation_candidate"):
                continue
            lines.append(
                f"{item['domain']}: resultado classificado como {item['category']}, com T score {item['t_score']}."
            )
            break
    return lines


def _build_neopi_factor_paragraphs(bundle: dict[str, Any], draft_sections: dict[str, dict[str, Any]]) -> list[str]:
    neopi_section = draft_sections.get("neopi", {})
    summary_map = {
        str(item.get("domain") or "").strip(): str(item.get("summary") or "").strip()
        for item in neopi_section.get("neopi_factor_summaries", [])
        if isinstance(item, dict) and str(item.get("domain") or "").strip() and str(item.get("summary") or "").strip()
    }
    if not summary_map:
        return []

    paragraphs: list[str] = []
    for domain_name in NEOPI_DOMAIN_ORDER:
        summary = summary_map.get(domain_name)
        if not summary:
            continue
        friendly_text = rewrite_neopi_synthesis_text(summary, domain_name=domain_name)
        if friendly_text:
            paragraphs.append(f"{domain_name}: {friendly_text}")
    return paragraphs


def _build_profiler_intro(bundle: dict[str, Any]) -> list[str]:
    dominant_style = bundle.get("profiler", {}).get("dominant_style")
    if not dominant_style:
        return []
    return [f"Neste momento, o estilo predominante identificado e {dominant_style}."]


def _build_section_paragraphs(
    section_key: str,
    bundle: dict[str, Any],
    draft_sections: dict[str, dict[str, Any]],
) -> list[str]:
    section = draft_sections.get(section_key, {})
    paragraphs = [str(item).strip() for item in section.get("paragraphs", []) if str(item).strip()]
    if paragraphs:
        return paragraphs

    raw_lines = bundle.get("report_template", {}).get("sections", {}).get(section_key, [])
    return [str(item).strip() for item in raw_lines if str(item).strip()]


def _build_conclusion_bullets(bundle: dict[str, Any], draft_sections: dict[str, dict[str, Any]]) -> list[str]:
    conclusion_section = draft_sections.get("conclusion", {})
    bullets = [str(item).strip() for item in conclusion_section.get("bullets", []) if str(item).strip()]
    if bullets:
        return [_strip_leading_numbering(item) for item in bullets]

    raw_lines = bundle.get("report_template", {}).get("sections", {}).get("conclusion", [])
    return [
        _strip_leading_numbering(str(item).strip())
        for item in raw_lines
        if str(item).strip() and "a partir dos indicadores" not in str(item).strip().lower()
    ]


def _build_procedures_label(bundle: dict[str, Any]) -> str:
    procedures = []
    if bundle.get("neopi"):
        procedures.append("Teste NEO PI-R")
    if bundle.get("profiler"):
        procedures.append("Profiler")
    if bundle.get("career_anchors"):
        procedures.append("Ancora de Carreira")
    if bundle.get("cultural_diagnosis"):
        procedures.append("Diagnostico Cultural")
    return ", ".join(procedures) if procedures else "Instrumentos comportamentais"


def _title_case(value: str) -> str:
    return value[:1].upper() + value[1:] if value else value


def _strip_leading_numbering(value: str) -> str:
    return re.sub(r"^\s*\d+\s*[-.)]?\s*", "", value).strip()
