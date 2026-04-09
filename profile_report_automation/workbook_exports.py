from __future__ import annotations

from copy import copy
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from profile_report_automation.neopi_language import (
    NEOPI_DISPLAY_LABELS,
    NEOPI_DOMAIN_ORDER,
    rewrite_neopi_synthesis_text,
)

ANCHOR_ROW_MAP = {
    "Técnico Funcional": 3,
    "Administrativo Geral": 4,
    "Autonomia Independência": 5,
    "Segurança Estabilidade": 6,
    "Criatividade Empreendedora": 7,
    "Vontade de Servir": 8,
    "Puro Desafio": 9,
    "Estilo de Vida": 10,
}

CULTURE_ORDER = ["Clã", "Inovativa", "Mercado", "Hierárquica"]

STATIC_SIGNATURE_PATTERNS = [
    re.compile(r"carolina de oliveira giarola", flags=re.IGNORECASE),
    re.compile(r"crp\s*-\s*04\s*/\s*81762", flags=re.IGNORECASE),
]


def export_filled_workbook(
    output_dir: str | Path,
    bundle: dict[str, Any],
    coverage: dict[str, Any],
    draft: dict[str, Any],
) -> str:
    template_path = _resolve_template_path(bundle)
    if not template_path:
        raise FileNotFoundError(
            "Nenhum modelo de planilha foi encontrado. Envie a planilha modelo uma vez para o sistema armazenar o template."
        )

    workbook = load_workbook(template_path)

    _fill_summary_sheet(workbook, bundle, draft)
    _fill_profiler_sheet(workbook, bundle)
    _fill_anchor_sheet(workbook, bundle)
    _fill_culture_sheet(workbook, bundle)
    _fill_reference_sheet(workbook, bundle)
    _remove_internal_review_sheet(workbook)
    _remove_static_signatures(workbook)

    output_root = Path(output_dir)
    output_root.mkdir(parents=True, exist_ok=True)
    output_path = output_root / f"{_build_report_basename(bundle)}.xlsx"
    workbook.save(output_path)
    return str(output_path.resolve())


def _resolve_template_path(bundle: dict[str, Any]) -> Path | None:
    raw_path = str(bundle.get("files", {}).get("report_workbook", "") or "").strip()
    if raw_path:
        direct_path = Path(raw_path).expanduser()
        if direct_path.exists():
            return direct_path
    return _find_cached_template_path()


def _find_cached_template_path() -> Path | None:
    repo_root = Path(__file__).resolve().parents[1]
    templates_root = repo_root / "artifacts" / "templates"
    uploads_root = repo_root / "artifacts" / "uploads"

    template_candidates = [
        path
        for path in templates_root.rglob("*.xlsx")
        if templates_root.exists() and _is_report_template_candidate(path)
    ]
    if template_candidates:
        return max(template_candidates, key=lambda item: item.stat().st_mtime)

    upload_candidates = [
        path
        for path in uploads_root.rglob("*.xlsx")
        if uploads_root.exists() and path.parent.name == "intake" and _is_report_template_candidate(path)
    ]
    if upload_candidates:
        return max(upload_candidates, key=lambda item: item.stat().st_mtime)

    return None


def _is_report_template_candidate(path: Path) -> bool:
    normalized = _normalize_lookup_text(path.name)
    return normalized.endswith(".xlsx") and "relatorio de analise de perfil" in normalized


def _normalize_lookup_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    without_marks = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", without_marks).strip().lower()


def _format_application_date(value: Any) -> str:
    raw_value = str(value or "").strip()
    if not raw_value:
        return "-"

    for date_format in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw_value, date_format).strftime("%d/%m/%Y")
        except ValueError:
            continue

    return raw_value


def _format_generation_date() -> str:
    return datetime.now().strftime("%d/%m/%Y")


def _remove_static_signatures(workbook: Any) -> None:
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                cleaned_value = _strip_static_signature_lines(cell.value)
                if cleaned_value != cell.value:
                    cell.value = cleaned_value


def _strip_static_signature_lines(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    kept_lines = [
        line
        for line in lines
        if line and not any(pattern.search(line) for pattern in STATIC_SIGNATURE_PATTERNS)
    ]
    return "\n".join(kept_lines)


def _fill_summary_sheet(workbook: Any, bundle: dict[str, Any], draft: dict[str, Any]) -> None:
    ws = workbook["Síntese"]
    person = bundle.get("person", {})
    ws["B7"] = f"Nome: {person.get('name') or '-'}"
    ws["B8"] = f"Data da aplicação: {_format_application_date(person.get('application_date'))}"
    ws["B9"] = person.get("business_unit") or "-"
    ws["B10"] = f"Demanda: {person.get('demand') or '-'}"
    ws["B11"] = f"Procedimentos realizados: {_build_procedures_label(bundle)}"

    neopi_lines = _build_neopi_summary_lines(bundle, draft)
    for row, text in zip(range(21, 26), neopi_lines):
        ws[f"B{row}"] = text

    dominant_style = str(bundle.get("profiler", {}).get("dominant_style") or "-")
    ws["B37"] = f"Nesse momento, apresenta o estilo: {dominant_style}."
    ws["B38"] = _build_profiler_summary(bundle)

    anchor_rows = _build_anchor_summary_lines(bundle)
    for row, text in zip([52, 53], anchor_rows):
        ws[f"B{row}"] = text
    for row in [52, 53]:
        if row - 52 >= len(anchor_rows):
            ws[f"B{row}"] = ""

    culture_rows = _build_culture_summary_lines(bundle)
    for row, text in zip([71, 72], culture_rows):
        ws[f"B{row}"] = text
    for row in [71, 72]:
        if row - 71 >= len(culture_rows):
            ws[f"B{row}"] = ""

    conclusion_lines = _build_conclusion_lines(bundle, draft)
    ws["B74"] = "A partir dos indicadores de seu perfil, apresenta:"
    for index, row in enumerate(range(75, 80), start=1):
        if index <= len(conclusion_lines):
            ws[f"B{row}"] = f"{index}- {conclusion_lines[index - 1]}"
        else:
            ws[f"B{row}"] = ""


def _fill_profiler_sheet(workbook: Any, bundle: dict[str, Any]) -> None:
    ws = workbook["Input Profiler"]
    score_map = {str(item["style"]): float(item["score"]) for item in bundle.get("profiler", {}).get("scores", [])}
    for row in range(2, 6):
        style = str(ws[f"A{row}"].value or "").strip()
        if style in score_map:
            ws[f"B{row}"] = score_map[style]


def _fill_anchor_sheet(workbook: Any, bundle: dict[str, Any]) -> None:
    ws = workbook["Imput âncora de Carreira"]
    anchor_map = {str(item["name"]): item for item in bundle.get("career_anchors", {}).get("scores", [])}

    for anchor_name, row in ANCHOR_ROW_MAP.items():
        item = anchor_map.get(anchor_name)
        ws[f"D{row}"] = _format_anchor_label(anchor_name)
        if item:
            description = str(item.get("description") or "").strip()
            prefix = anchor_name.replace("Independência", "e Independência")
            if not str(ws[f"A{row}"].value or "").strip():
                ws[f"A{row}"] = f"{prefix}: {description}" if description else prefix
            ws[f"E{row}"] = float(item["average"])
        else:
            if not str(ws[f"A{row}"].value or "").strip():
                ws[f"A{row}"] = anchor_name
            ws[f"E{row}"] = ""


def _fill_culture_sheet(workbook: Any, bundle: dict[str, Any]) -> None:
    ws = workbook["Input Cultura Organizacional"]
    culture_map = {str(item["culture"]): item for item in bundle.get("cultural_diagnosis", {}).get("scores", [])}

    ws["E2"] = "Minha preferência"
    for index, culture_name in enumerate(CULTURE_ORDER, start=2):
        item = culture_map.get(culture_name, {})
        description = str(item.get("description") or "").strip()
        if description and not str(ws[f"B{index}"].value or "").strip():
            ws[f"B{index}"] = f"{culture_name}: {description}"

    for index, culture_name in enumerate(CULTURE_ORDER, start=3):
        item = culture_map.get(culture_name, {})
        ws[f"D{index}"] = culture_name
        ws[f"E{index}"] = round(float(item.get("score", 0)), 2) if item else ""


def _fill_reference_sheet(workbook: Any, bundle: dict[str, Any]) -> None:
    ws = workbook["Dados gerais"]
    intros = {
        "B3": "O Inventário de Personalidade NEO Revisado (NEO PI-R) é um instrumento de avaliação da personalidade, validado pelo CFP, baseado no modelo clássico dos Cinco Grandes Fatores, que compreendem a personalidade a partir da combinação de 5 grandes fatores.",
        "B4": "Profiler é uma ferramenta para mapeamento de estilo comportamental baseado em quatro características comportamentais: Executores, Comunicadores, Planejadores e Analistas.",
        "B5": "O questionário visa a reflexão sobre o que é mais importante para a pessoa em suas áreas de competência, objetivos e valores, a partir da categorização em 8 âncoras de carreira.",
        "B6": "O questionário adaptado da teoria Competing Values identifica as preferências de cultura organizacional em quatro tipos: hierárquica, clã, inovadora e de mercado.",
    }

    for cell, text in intros.items():
        if not str(ws[cell].value or "").strip():
            ws[cell] = text


def _fill_neo_review_sheet(workbook: Any, bundle: dict[str, Any], coverage: dict[str, Any]) -> None:
    title = "Base Automacao NEO PI-R"
    if title in workbook.sheetnames:
        del workbook[title]

    ws = workbook.create_sheet(title)
    ws["A1"] = "Base interna de revisão do NEO PI-R"
    ws["A2"] = "Esses itens apoiam a revisão e não precisam entrar automaticamente no relatório final."
    ws.append([])
    ws.append(["Tipo", "Nome", "Domínio", "Categoria", "T Score", "Entrar no relatório final?"])

    for item in bundle.get("neopi", {}).get("domains", []):
        ws.append(
            [
                "Fator",
                item.get("domain"),
                item.get("domain"),
                item.get("category"),
                item.get("t_score"),
                "Sim" if item.get("citation_candidate") else "Opcional",
            ]
        )

    for item in bundle.get("neopi", {}).get("facets", []):
        if not item.get("citation_candidate"):
            continue
        ws.append(
            [
                "Faceta",
                item.get("facet"),
                item.get("domain"),
                item.get("category"),
                item.get("t_score"),
                "Revisar",
            ]
        )

    ws.append([])
    ws.append(["Sinais obrigatórios mapeados automaticamente"])
    ws.append(["Signal ID", "Nome", "Fonte", "Coberto no relatório", "Coberto na conclusão"])
    for signal in coverage.get("required_signals", []):
        if signal.get("source_section") != "neopi":
            continue
        ws.append(
            [
                signal.get("signal_id"),
                signal.get("name"),
                signal.get("source_section"),
                signal.get("coverage", {}).get("all_sections", {}).get("covered"),
                signal.get("coverage", {}).get("conclusion", {}).get("covered"),
            ]
        )


def _remove_internal_review_sheet(workbook: Any) -> None:
    title = "Base Automacao NEO PI-R"
    if title in workbook.sheetnames:
        del workbook[title]


def _build_neopi_summary_lines(bundle: dict[str, Any], draft: dict[str, Any]) -> list[str]:
    synthesis_map = _draft_neopi_summary_map(draft) or bundle.get("neopi", {}).get("friendly_synthesis_by_domain") or bundle.get("neopi", {}).get("synthesis_by_domain", {})
    lines: list[str] = []

    for domain_name in NEOPI_DOMAIN_ORDER:
        raw_text = str(synthesis_map.get(domain_name) or "").strip()
        cleaned_text = rewrite_neopi_synthesis_text(raw_text, domain_name=domain_name)
        display_name = NEOPI_DISPLAY_LABELS.get(domain_name, domain_name)
        if cleaned_text:
            lines.append(f"{display_name}: {cleaned_text}")
        else:
            fallback = _build_neopi_domain_fallback(bundle, domain_name)
            lines.append(f"{display_name}: {fallback}")

    return lines


def _draft_neopi_summary_map(draft: dict[str, Any]) -> dict[str, str]:
    summary_map: dict[str, str] = {}
    for item in draft.get("neopi_factor_summaries", []):
        if not isinstance(item, dict):
            continue
        domain_name = str(item.get("domain") or "").strip()
        summary = str(item.get("summary") or "").strip()
        if domain_name and summary:
            summary_map[domain_name] = summary
    return summary_map


def _build_neopi_domain_fallback(bundle: dict[str, Any], domain_name: str) -> str:
    for item in bundle.get("neopi", {}).get("domains", []):
        if item.get("domain") == domain_name:
            return f"Resultado classificado como {item.get('category')} no momento atual, com T score {item.get('t_score')}."
    return "Sem conteúdo identificado para este fator."


def _build_profiler_summary(bundle: dict[str, Any]) -> str:
    style = str(bundle.get("profiler", {}).get("dominant_style") or "").strip()
    summaries = {
        "Executor": "Apresenta energia para ação, decisão e enfrentamento de desafios, com tendência a imprimir ritmo, objetividade e senso de urgência ao trabalho.",
        "Comunicador": "Tende a se comunicar com fluidez, fortalecer relações e mobilizar pessoas com entusiasmo, abertura e facilidade de interação.",
        "Planejador": "Costuma atuar com prudência, observação e consistência, valorizando previsibilidade, organização e estabilidade na execução.",
        "Analista": "Tende a priorizar precisão, profundidade e senso crítico, com atenção elevada a detalhes, critérios e qualidade das entregas.",
    }
    return summaries.get(style, "O estilo predominante identificado contribui para a leitura do comportamento no contexto profissional.")


def _build_anchor_summary_lines(bundle: dict[str, Any]) -> list[str]:
    items = bundle.get("career_anchors", {}).get("top_anchors", [])
    lines = []
    for item in items[:2]:
        name = str(item.get("name") or "").strip()
        description = str(item.get("description") or "").strip()
        if name:
            lines.append(f"{name}: {description}" if description else name)
    return lines


def _build_culture_summary_lines(bundle: dict[str, Any]) -> list[str]:
    items = bundle.get("cultural_diagnosis", {}).get("top_cultures", [])
    lines = []
    for item in items[:2]:
        name = str(item.get("culture") or "").strip()
        description = str(item.get("description") or "").strip()
        if name:
            lines.append(f"{name}: {description}" if description else name)
    return lines


def _build_conclusion_lines(bundle: dict[str, Any], draft: dict[str, Any]) -> list[str]:
    for section in draft.get("sections", []):
        if section.get("key") == "conclusion":
            bullets = [str(item).strip() for item in section.get("bullets", []) if str(item).strip()]
            if bullets:
                return [_strip_numbering(item) for item in bullets][:5]

    raw_lines = bundle.get("report_template", {}).get("sections", {}).get("conclusion", [])
    cleaned_lines = [
        _strip_numbering(str(item).strip())
        for item in raw_lines
        if str(item).strip() and "a partir dos indicadores" not in str(item).strip().lower()
    ]
    return cleaned_lines[:5]


def _build_procedures_label(bundle: dict[str, Any]) -> str:
    procedures = []
    if bundle.get("neopi"):
        procedures.append("Teste NEO PI-R")
    if bundle.get("profiler"):
        procedures.append("Profiler")
    if bundle.get("career_anchors"):
        procedures.append("Âncora de Carreira")
    if bundle.get("cultural_diagnosis"):
        procedures.append("Diagnóstico Cultural")
    return ", ".join(procedures)


def _format_anchor_label(anchor_name: str) -> str:
    if anchor_name == "Autonomia Independência":
        return "Autonomia\nIndependência"
    if anchor_name == "Segurança Estabilidade":
        return "Segurança \nEstabilidade"
    if anchor_name == "Criatividade Empreendedora":
        return "Criatividade\nEmpreendedora"
    return anchor_name


def _strip_numbering(value: str) -> str:
    return re.sub(r"^\s*\d+\s*[-.)]?\s*", "", value).strip()


def _build_report_basename(bundle: dict[str, Any]) -> str:
    raw_name = str(bundle.get("person", {}).get("name") or "Pessoa avaliada")
    normalized = unicodedata.normalize("NFKD", raw_name)
    without_marks = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    safe_name = re.sub(r'[<>:"/\\\\|?*]+', " ", without_marks)
    safe_name = re.sub(r"\s+", " ", safe_name).strip() or "Pessoa avaliada"
    return f"{safe_name} - Relatorio de Analise de Perfil"
