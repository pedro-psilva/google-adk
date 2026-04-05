#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import load_workbook
from pypdf import PdfReader

NEOPI_DOMAINS = [
    "Neuroticismo",
    "Extroversão",
    "Abertura",
    "Amabilidade",
    "Conscienciosidade",
]

NEOPI_FACETS = {
    "Neuroticismo": [
        "Ansiedade",
        "Raiva",
        "Depressão",
        "Embaraço",
        "Impulsividade",
        "Vulnerabilidade",
    ],
    "Extroversão": [
        "Acolhimento caloroso",
        "Gregarismo",
        "Assertividade",
        "Atividade",
        "Busca de sensações",
        "Emoções positivas",
    ],
    "Abertura": [
        "Fantasia",
        "Estética",
        "Sentimentos",
        "Ações variadas",
        "Ideias",
        "Valores",
    ],
    "Amabilidade": [
        "Confiança",
        "Franqueza",
        "Altruísmo",
        "Complacência",
        "Modéstia",
        "Sensibilidade",
    ],
    "Conscienciosidade": [
        "Competência",
        "Ordem",
        "Senso de dever",
        "Esforço por realizações",
        "Autodisciplina",
        "Ponderação",
    ],
}

CAREER_ANCHOR_DESCRIPTIONS = {
    "Técnico Funcional": (
        "Concentra-se no desenvolvimento de expertise e especialização em uma área "
        "específica. Busca ser especialista em seu campo de atuação, valorizando o "
        "reconhecimento entre outros especialistas ou colegas."
    ),
    "Administrativo Geral": (
        "Destaca-se pela integração dos esforços coletivos para alcançar resultados "
        "e pela coordenação de diferentes funções em uma organização."
    ),
    "Autonomia Independência": (
        "Valoriza a liberdade de agir conforme suas próprias normas, buscando atuar "
        "de maneira mais independente e flexível."
    ),
    "Segurança Estabilidade": (
        "Busca tranquilidade, segurança financeira e estabilidade no trabalho."
    ),
    "Criatividade Empreendedora": (
        "Valoriza criar algo novo, superar obstáculos e influenciar outros por meio "
        "de ideias e empreendimentos."
    ),
    "Vontade de Servir": (
        "Usa habilidades interpessoais em prol de uma causa significativa e valoriza "
        "contribuir para a melhoria social."
    ),
    "Puro Desafio": (
        "Tem foco na superação de desafios difíceis e na resolução de problemas "
        "aparentemente insolúveis."
    ),
    "Estilo de Vida": (
        "Busca integrar carreira, família e necessidades pessoais, valorizando "
        "flexibilidade e equilíbrio entre essas esferas."
    ),
}

CULTURE_DESCRIPTIONS = {
    "Clã": (
        "Valoriza uma cultura organizacional flexível orientada às relações humanas, "
        "cooperação e colaboração."
    ),
    "Inovativa": (
        "Valoriza uma cultura organizacional aberta e flexível, voltada a novas "
        "ideias, criatividade e adaptação."
    ),
    "Mercado": (
        "Valoriza uma cultura orientada à produtividade, competitividade e foco em "
        "metas e resultados."
    ),
    "Hierárquica": (
        "Valoriza ordem, previsibilidade, estabilidade, regras claras e processos "
        "estruturados."
    ),
}

REPORT_FILE_PATTERNS = {
    "report_pdf": ["*Relatório de Análise de Perfil.pdf", "*Relatorio de Analise de Perfil.pdf"],
    "report_workbook": ["*Relatório de Análise de Perfil.xlsx", "*Relatorio de Analise de Perfil.xlsx"],
    "neopi_pdf": ["*NEOPI-R*pdf"],
    "profiler_pdf": ["*extended.pdf"],
    "anchors_workbook": ["*IEBT Innovation.xlsx"],
}


def normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    without_marks = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", without_marks).strip().lower()


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\x00", " ")).strip()


def classify_t_score(t_score: int) -> str:
    if t_score <= 34:
        return "muito baixo"
    if t_score <= 44:
        return "baixo"
    if t_score <= 55:
        return "medio"
    if t_score <= 65:
        return "alto"
    return "muito alto"


def is_extreme(category: str) -> bool:
    return category != "medio"


def find_first_file(base_dir: Path, patterns: list[str]) -> Path | None:
    matches: list[Path] = []
    seen: set[str] = set()
    for pattern in patterns:
        for path in sorted(base_dir.glob(pattern)):
            resolved = str(path.resolve())
            if resolved not in seen:
                seen.add(resolved)
                matches.append(path)
    return matches[0] if matches else None


def extract_pdf_pages(path: Path, engine: str = "pypdf") -> list[str]:
    if engine == "pdfplumber":
        with pdfplumber.open(path) as pdf:
            return [clean_text(page.extract_text() or "") for page in pdf.pages]

    reader = PdfReader(str(path))
    return [clean_text(page.extract_text() or "") for page in reader.pages]


def parse_score_from_text(text: str, label: str) -> dict[str, int] | None:
    normalized_text = normalize_text(text)
    normalized_label = normalize_text(label)
    pattern = re.compile(rf"\b{re.escape(normalized_label)}\b\s+(\d+)\s+(\d+)\b")
    match = pattern.search(normalized_text)
    if not match:
        return None
    return {
        "raw_score": int(match.group(1)),
        "t_score": int(match.group(2)),
    }


def extract_neopi_synthesis(page_text: str) -> dict[str, str]:
    headings = [
        "NEUROTICISMO",
        "EXTROVERSÃO",
        "ABERTURA",
        "AMABILIDADE",
        "CONSCIENCIOSIDADE",
    ]
    sections: dict[str, str] = {}
    for index, heading in enumerate(headings):
        start = page_text.find(heading)
        if start == -1:
            continue
        content_start = start + len(heading)
        end = len(page_text)
        for next_heading in headings[index + 1 :]:
            next_pos = page_text.find(next_heading, content_start)
            if next_pos != -1:
                end = next_pos
                break
        sections[heading.title()] = clean_text(page_text[content_start:end])
    return sections


def extract_neopi(path: Path) -> dict[str, Any]:
    pages = extract_pdf_pages(path, engine="pypdf")
    score_page = pages[3] if len(pages) > 3 else ""
    synthesis_page = pages[10] if len(pages) > 10 else ""

    domains: list[dict[str, Any]] = []
    facets: list[dict[str, Any]] = []

    for domain_name in NEOPI_DOMAINS:
        scores = parse_score_from_text(score_page, domain_name)
        if not scores:
            continue
        category = classify_t_score(scores["t_score"])
        domains.append(
            {
                "domain": domain_name,
                **scores,
                "category": category,
                "citation_candidate": is_extreme(category),
            }
        )

    for domain_name, facet_names in NEOPI_FACETS.items():
        for facet_name in facet_names:
            scores = parse_score_from_text(score_page, facet_name)
            if not scores:
                continue
            category = classify_t_score(scores["t_score"])
            facets.append(
                {
                    "domain": domain_name,
                    "facet": facet_name,
                    **scores,
                    "category": category,
                    "citation_candidate": is_extreme(category),
                }
            )

    citation_candidates = [
        *[
            {
                "type": "domain",
                "name": domain["domain"],
                "category": domain["category"],
                "t_score": domain["t_score"],
            }
            for domain in domains
            if domain["citation_candidate"]
        ],
        *[
            {
                "type": "facet",
                "name": facet["facet"],
                "domain": facet["domain"],
                "category": facet["category"],
                "t_score": facet["t_score"],
            }
            for facet in facets
            if facet["citation_candidate"]
        ],
    ]

    return {
        "pages": len(pages),
        "domains": domains,
        "facets": facets,
        "citation_candidates": citation_candidates,
        "synthesis_by_domain": extract_neopi_synthesis(synthesis_page),
    }


def extract_prefixed_value(ws: Any, prefix: str) -> str | None:
    normalized_prefix = normalize_text(prefix)
    for row in ws.iter_rows():
        for cell in row:
            value = clean_text(cell.value)
            if not value:
                continue
            normalized_value = normalize_text(value)
            if normalized_value.startswith(normalized_prefix):
                parts = value.split(":", 1)
                if len(parts) == 2:
                    return clean_text(parts[1])
                if value.startswith(prefix):
                    return clean_text(value[len(prefix) :])
                return value
    return None


def find_row_in_column(ws: Any, column: str, marker: str) -> int | None:
    normalized_marker = normalize_text(marker)
    for row in range(1, ws.max_row + 1):
        value = clean_text(ws[f"{column}{row}"].value)
        if value and normalized_marker in normalize_text(value):
            return row
    return None


def collect_column_section(ws: Any, start_marker: str, end_markers: list[str]) -> list[str]:
    start_row = find_row_in_column(ws, "B", start_marker)
    if start_row is None:
        return []

    end_row = ws.max_row + 1
    for marker in end_markers:
        marker_row = find_row_in_column(ws, "B", marker)
        if marker_row and marker_row > start_row:
            end_row = min(end_row, marker_row)

    values: list[str] = []
    for row in range(start_row + 1, end_row):
        value = clean_text(ws[f"B{row}"].value)
        if not value:
            continue
        if normalize_text(value).startswith("pagina "):
            continue
        values.append(value)
    return values


def extract_profiler_scores(ws: Any) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        style = clean_text(ws[f"A{row}"].value)
        score = ws[f"B{row}"].value
        if not style or score is None:
            continue
        percentage = round(float(score) * 100, 2)
        items.append(
            {
                "style": style,
                "score": round(float(score), 4),
                "percentage": percentage,
            }
        )
    return items


def canonical_anchor_name(name: str) -> str:
    simplified = clean_text(name).replace("\n", " ")
    simplified = re.sub(r"\s+", " ", simplified).strip()
    aliases = {
        "Administração geral": "Administrativo Geral",
        "Autonomia e Independência": "Autonomia Independência",
        "Segurança e Estabilidade": "Segurança Estabilidade",
        "Técnico Funcional:": "Técnico Funcional",
        "Administração geral:": "Administrativo Geral",
        "Autonomia e Independência:": "Autonomia Independência",
        "Segurança e Estabilidade:": "Segurança Estabilidade",
        "Criatividade Empreendedora:": "Criatividade Empreendedora",
        "Vontade de servir:": "Vontade de Servir",
        "Puro desafio:": "Puro Desafio",
        "Estilo de Vida:": "Estilo de Vida",
        "Autonomia Independência": "Autonomia Independência",
        "Segurança Estabilidade": "Segurança Estabilidade",
        "Criatividade Empreendedora": "Criatividade Empreendedora",
        "Técnico Funcional": "Técnico Funcional",
        "Administrativo Geral": "Administrativo Geral",
        "Vontade de Servir": "Vontade de Servir",
        "Puro Desafio": "Puro Desafio",
        "Estilo de Vida": "Estilo de Vida",
    }
    normalized = normalize_text(simplified)
    for candidate in aliases:
        if normalize_text(candidate) == normalized:
            return aliases[candidate]
    return simplified


def extract_report_anchor_notes(ws: Any) -> list[str]:
    notes: list[str] = []
    mismatches: list[str] = []
    for row in range(3, 11):
        description = clean_text(ws[f"A{row}"].value)
        anchor_name = canonical_anchor_name(clean_text(ws[f"D{row}"].value))
        if not description or not anchor_name:
            continue
        description_anchor = canonical_anchor_name(description.split(":", 1)[0])
        if normalize_text(description_anchor) != normalize_text(anchor_name):
            mismatches.append(f"row {row}: descricao={description_anchor} | rotulo={anchor_name}")
    if mismatches:
        notes.append(
            "Career anchor description rows do not fully align with anchor labels in the "
            "report workbook input sheet: " + "; ".join(mismatches)
        )
    return notes


def extract_report_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True)
    summary_ws = workbook.worksheets[0]
    profiler_ws = workbook.worksheets[2]
    anchor_input_ws = workbook.worksheets[3]
    culture_ws = workbook.worksheets[4]

    sections = {
        "neopi": collect_column_section(summary_ws, "3.1- NEOPI-R", ["3.2- Profiler"]),
        "profiler": collect_column_section(summary_ws, "3.2- Profiler", ["3.3- Âncoras de Carreira"]),
        "career_anchors": collect_column_section(
            summary_ws,
            "3.3- Âncoras de Carreira",
            ["3.4- Diagnóstico Cultural"],
        ),
        "cultural_diagnosis": collect_column_section(
            summary_ws,
            "3.4- Diagnóstico Cultural",
            ["IV. Conclusão"],
        ),
        "conclusion": collect_column_section(summary_ws, "IV. Conclusão", []),
    }

    profiler_scores = extract_profiler_scores(profiler_ws)
    dominant_profiler = max(profiler_scores, key=lambda item: item["percentage"]) if profiler_scores else None

    culture_scores: list[dict[str, Any]] = []
    for row in range(3, 7):
        culture_name = clean_text(culture_ws[f"D{row}"].value)
        score = culture_ws[f"E{row}"].value
        if not culture_name or score is None:
            continue
        culture_scores.append(
            {
                "culture": culture_name,
                "score": float(score),
                "description": CULTURE_DESCRIPTIONS.get(culture_name, ""),
            }
        )

    return {
        "person": {
            "name": extract_prefixed_value(summary_ws, "Nome:"),
            "application_date": extract_prefixed_value(summary_ws, "Data da aplicação:"),
            "business_unit": extract_prefixed_value(summary_ws, "Unidade de negócios"),
            "demand": extract_prefixed_value(summary_ws, "Demanda:"),
        },
        "profiler_scores": profiler_scores,
        "dominant_profiler_style": dominant_profiler["style"] if dominant_profiler else None,
        "cultural_scores_from_template": culture_scores,
        "sections": sections,
        "notes": extract_report_anchor_notes(anchor_input_ws),
    }


def extract_anchor_summary(ws: Any) -> list[dict[str, Any]]:
    start_row = None
    for row in range(1, ws.max_row + 1):
        values = [clean_text(cell.value) for cell in ws[row]]
        if any(normalize_text(value) == normalize_text("Âncoras") for value in values if value):
            start_row = row
            break
    if start_row is None:
        return []

    items: list[dict[str, Any]] = []
    for row in range(start_row + 1, start_row + 9):
        name = canonical_anchor_name(clean_text(ws[f"B{row}"].value))
        average = ws[f"D{row}"].value
        if not name or average is None:
            continue
        items.append(
            {
                "name": name,
                "average": float(average),
                "description": CAREER_ANCHOR_DESCRIPTIONS.get(name, ""),
            }
        )
    return items


def extract_culture_summary(ws: Any) -> list[dict[str, Any]]:
    summary_row = None
    for row in range(1, ws.max_row + 1):
        values = [clean_text(cell.value) for cell in ws[row]]
        if any(normalize_text(value) == normalize_text("Minha preferência") for value in values if value):
            summary_row = row
    if summary_row is None:
        return []

    items: list[dict[str, Any]] = []
    for row in range(summary_row + 1, summary_row + 5):
        name = clean_text(ws[f"B{row}"].value)
        score = ws[f"C{row}"].value
        if not name or score is None:
            continue
        canonical_name = canonical_anchor_name(name)
        items.append(
            {
                "culture": canonical_name,
                "score": round(float(score), 4),
                "description": CULTURE_DESCRIPTIONS.get(canonical_name, ""),
            }
        )
    return items


def extract_anchor_and_culture_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True)
    anchors_ws = workbook.worksheets[0]
    culture_ws = workbook.worksheets[1]

    anchor_scores = extract_anchor_summary(anchors_ws)
    top_anchors = sorted(anchor_scores, key=lambda item: item["average"], reverse=True)[:2]

    culture_scores = extract_culture_summary(culture_ws)
    top_cultures = sorted(culture_scores, key=lambda item: item["score"], reverse=True)[:2]

    return {
        "person": {
            "name": clean_text(anchors_ws["C2"].value),
            "role": clean_text(anchors_ws["C3"].value),
            "application_date": clean_text(anchors_ws["C4"].value),
        },
        "career_anchors": {
            "scores": anchor_scores,
            "top_anchors": top_anchors,
        },
        "cultural_diagnosis": {
            "scores": culture_scores,
            "top_cultures": top_cultures,
        },
    }


def extract_profiler_extended(path: Path) -> dict[str, Any]:
    pages = extract_pdf_pages(path, engine="pdfplumber")
    page_two = pages[1] if len(pages) > 1 else ""
    page_seven = pages[6] if len(pages) > 6 else ""
    page_eight = pages[7] if len(pages) > 7 else ""

    dominant_match = re.search(r"Neste momento, .* está:\s*([A-Za-zÀ-ÿ]+)\s+em", page_two)
    dominant_style = dominant_match.group(1) if dominant_match else None

    long_form_sections = {
        "page_2_overview": page_two,
        "page_7_management": page_seven,
        "page_8_sales_and_motivation": page_eight,
    }

    return {
        "pages": len(pages),
        "dominant_style_from_pdf": dominant_style,
        "long_form_sections": long_form_sections,
    }


def build_bundle(base_dir: Path) -> dict[str, Any]:
    files = {key: find_first_file(base_dir, patterns) for key, patterns in REPORT_FILE_PATTERNS.items()}
    missing = [key for key, value in files.items() if value is None]
    if missing:
        raise FileNotFoundError(f"Missing expected files for keys: {', '.join(missing)}")

    report_workbook = extract_report_workbook(files["report_workbook"])
    anchor_bundle = extract_anchor_and_culture_workbook(files["anchors_workbook"])

    notes = []
    notes.extend(report_workbook.get("notes", []))
    if report_workbook["dominant_profiler_style"] and anchor_bundle["career_anchors"]["top_anchors"]:
        notes.append(
            "The sample suggests a concise final report can be built from a small subset "
            "of signals: dominant profiler style, top-2 anchors, top-2 cultures, and NEO "
            "non-medium indicators."
        )

    person = {
        "name": report_workbook["person"].get("name") or anchor_bundle["person"].get("name"),
        "application_date": report_workbook["person"].get("application_date")
        or anchor_bundle["person"].get("application_date"),
        "business_unit": report_workbook["person"].get("business_unit"),
        "demand": report_workbook["person"].get("demand"),
        "role": anchor_bundle["person"].get("role"),
    }

    return {
        "input_dir": str(base_dir.resolve()),
        "files": {key: str(value.resolve()) for key, value in files.items()},
        "person": person,
        "neopi": extract_neopi(files["neopi_pdf"]),
        "profiler": {
            "scores": report_workbook["profiler_scores"],
            "dominant_style": report_workbook["dominant_profiler_style"],
            "extended_pdf": extract_profiler_extended(files["profiler_pdf"]),
        },
        "career_anchors": anchor_bundle["career_anchors"],
        "cultural_diagnosis": anchor_bundle["cultural_diagnosis"],
        "report_template": {
            "sections": report_workbook["sections"],
        },
        "notes": notes,
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extract a profile-report bundle into a normalized JSON payload."
    )
    parser.add_argument("input_dir", help="Folder containing the sample PDFs and spreadsheets")
    parser.add_argument(
        "--output",
        help="Optional JSON output path. If omitted, the payload is printed to stdout.",
    )
    args = parser.parse_args()

    bundle = build_bundle(Path(args.input_dir))
    payload = json.dumps(bundle, ensure_ascii=False, indent=2)

    if args.output:
        output_path = Path(args.output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(payload, encoding="utf-8")
        print(f"Wrote normalized bundle to {output_path.resolve()}")
        return

    print(payload)


if __name__ == "__main__":
    main()
