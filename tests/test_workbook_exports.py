from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils.cell import range_boundaries


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from profile_report_automation.workbook_exports import _compact_worksheet_layout, _fit_summary_sheet_text_rows  # noqa: E402


class WorkbookSummaryLayoutTests(unittest.TestCase):
    def test_fit_summary_sheet_text_rows_expands_rows_for_wrapped_content(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Síntese"
        worksheet.column_dimensions["B"].width = 87.88
        worksheet["B71"] = (
            "Clã: Valoriza uma cultura organizacional flexível orientada às relações humanas, "
            "cooperação e colaboração."
        )
        worksheet["B75"] = (
            "1- Monitorar riscos de sobrecarga e dispersão quando houver excesso de demandas "
            "simultâneas e mudanças frequentes de prioridade."
        )

        _fit_summary_sheet_text_rows(worksheet)

        self.assertGreater(worksheet.row_dimensions[71].height or 0, 15.75)
        self.assertGreater(worksheet.row_dimensions[75].height or 0, 15.75)
        self.assertTrue(worksheet["B71"].alignment.wrap_text)
        self.assertEqual(worksheet["B71"].alignment.vertical, "top")

    def test_compact_summary_layout_keeps_chart_columns_in_print_area(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "summary-with-chart.xlsx"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Síntese"
            worksheet["B21"] = "Texto de síntese"

            data_sheet = workbook.create_sheet("Dados")
            for index, value in enumerate([0.42, 0.66, 0.81], start=1):
                data_sheet[f"A{index}"] = f"Indicador {index}"
                data_sheet[f"B{index}"] = value

            chart = BarChart()
            chart.width = 15
            chart.height = 7.5
            chart.add_data(Reference(data_sheet, min_col=2, min_row=1, max_row=3))
            chart.set_categories(Reference(data_sheet, min_col=1, min_row=1, max_row=3))
            worksheet.add_chart(chart, "B32")
            workbook.save(workbook_path)

            reloaded_workbook = load_workbook(workbook_path)
            reloaded_worksheet = reloaded_workbook["Síntese"]

            _compact_worksheet_layout(reloaded_worksheet)

            print_area_range = str(reloaded_worksheet.print_area).split("!", maxsplit=1)[-1]
            _, _, max_column, _ = range_boundaries(print_area_range)
            self.assertGreater(max_column, 2)


if __name__ == "__main__":
    unittest.main()
