from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import load_workbook
from pypdf import PdfReader

from profile_report_automation.neopi_language import (
    NEOPI_DOMAIN_ORDER,
    build_friendly_neopi_synthesis_map,
    rewrite_neopi_synthesis_text,
)

NEOPI_DOMAINS = list(NEOPI_DOMAIN_ORDER)

NEOPI_SYNTHESIS_HEADINGS = [
    ("Neuroticismo", ["NEUROTICISMO"]),
    ("Extroversão", ["EXTROVERSÃO", "EXTROVERSAO", "EXTROVERSÃƒO"]),
    ("Abertura", ["ABERTURA"]),
    ("Amabilidade", ["AMABILIDADE"]),
    ("Conscienciosidade", ["CONSCIENCIOSIDADE"]),
]

NEOPI_FACETS = {
    "Neuroticismo": [
        "Ansiedade",
        "Raiva",
        "Depressão",
        "Embaraço",
        "Impulsividade",
        "Vulnerabilidade",
    ],
    "Extroversão": [
        "Acolhimento caloroso",
        "Gregarismo",
        "Assertividade",
        "Atividade",
        "Busca de sensações",
        "Emoções positivas",
    ],
    "Abertura": [
        "Fantasia",
        "Estética",
        "Sentimentos",
        "Ações variadas",
        "Ideias",
        "Valores",
    ],
    "Amabilidade": [
        "Confiança",
        "Franqueza",
        "Altruísmo",
        "Complacência",
        "Modéstia",
        "Sensibilidade",
    ],
    "Conscienciosidade": [
        "Competência",
        "Ordem",
        "Senso de dever",
        "Esforço por realizações",
        "Autodisciplina",
        "Ponderação",
    ],
}

CAREER_ANCHOR_DESCRIPTIONS = {
    "Técnico Funcional": (
        "Concentra-se no desenvolvimento de expertise e especialização em uma área "
        "específica. Busca ser especialista em seu campo de atuação, valorizando o "
        "reconhecimento entre outros especialistas ou colegas."
    ),
    "Administrativo Geral": (
        "Destaca-se pela integração dos esforços coletivos para alcançar resultados "
        "e pela coordenação de diferentes funções em uma organização."
    ),
    "Autonomia Independência": (
        "Valoriza a liberdade de agir conforme suas próprias normas, buscando atuar "
        "de maneira mais independente e flexível."
    ),
    "Segurança Estabilidade": (
        "Busca tranquilidade, segurança financeira e estabilidade no trabalho."
    ),
    "Criatividade Empreendedora": (
        "Valoriza criar algo novo, superar obstáculos e influenciar outros por meio "
        "de ideias e empreendimentos."
    ),
    "Vontade de Servir": (
        "Usa habilidades interpessoais em prol de uma causa significativa e valoriza "
        "contribuir para a melhoria social."
    ),
    "Puro Desafio": (
        "Tem foco na superação de desafios difíceis e na resolução de problemas "
        "aparentemente insolúveis."
    ),
    "Estilo de Vida": (
        "Busca integrar carreira, família e necessidades pessoais, valorizando "
        "flexibilidade e equilíbrio entre essas esferas."
    ),
}

CULTURE_DESCRIPTIONS = {
    "Clã": (
        "Valoriza uma cultura organizacional flexível orientada às relações humanas, "
        "cooperação e colaboração."
    ),
    "Inovativa": (
        "Valoriza uma cultura organizacional aberta e flexível, voltada a novas "
        "ideias, criatividade e adaptação."
    ),
    "Mercado": (
        "Valoriza uma cultura orientada à produtividade, competitividade e foco em "
        "metas e resultados."
    ),
    "Hierárquica": (
        "Valoriza ordem, previsibilidade, estabilidade, regras claras e processos "
        "estruturados."
    ),
}

REPORT_FILE_PATTERNS = {
    "report_pdf": ["*Relatório de Análise de Perfil.pdf", "*Relatorio de Analise de Perfil.pdf"],
    "report_workbook": ["*Relatório de Análise de Perfil.xlsx", "*Relatorio de Analise de Perfil.xlsx"],
    "neopi_pdf": ["*NEOPI-R*.pdf", "*NEO PI-R*.pdf", "*NEOPI-R*pdf", "*NEO PI-R*pdf"],
    "profiler_pdf": ["*extended.pdf", "*regular.pdf", "*extended*pdf", "*regular*pdf", "*perfil*pdf"],
    "anchors_workbook": ["*IEBT Innovation.xlsx", "*Ancoras*.xlsx", "*Ã‚ncoras*.xlsx", "*Diagnostico*.xlsx"],
}

REQUIRED_FILE_KEYS = ("neopi_pdf", "profiler_pdf", "anchors_workbook")
UPLOAD_MANIFEST_FILENAME = "upload-manifest.json"

def normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    without_marks = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", without_marks).strip().lower()


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\x00", " ")).strip()


def infer_person_name_from_path(path: Path | None) -> str:
    if path is None:
        return ""

    stem = clean_text(path.stem.replace("_", " "))
    cleanup_patterns = [
        r"\s*-\s*neo\s*pi-?r.*$",
        r"\s*-\s*neopi-?r.*$",
        r"\s*-\s*regular.*$",
        r"\s*-\s*extended.*$",
        r"\s*-\s*relatorio de analise de perfil.*$",
        r"\s*-\s*\d{5,}.*$",
        r"\s*-?\s*\bprofiler\b.*$",
    ]
    candidate = stem
    for pattern in cleanup_patterns:
        candidate = re.sub(pattern, "", candidate, flags=re.IGNORECASE)

    candidate = candidate.replace("-", " ")
    candidate = re.sub(r"\s+", " ", candidate).strip(" -_")
    return candidate


def normalize_person_name_for_match(value: str) -> str:
    normalized = normalize_text(value)
    normalized = re.sub(r"[^a-z0-9\s]", " ", normalized)
    return re.sub(r"\s+", " ", normalized).strip()


def tokenize_person_name(value: str) -> list[str]:
    return [token for token in normalize_person_name_for_match(value).split(" ") if token]


def tokens_compatible(left: str, right: str) -> bool:
    if left == right:
        return True
    if len(left) == 1 and right.startswith(left):
        return True
    if len(right) == 1 and left.startswith(right):
        return True
    return False


def person_names_match(left: str, right: str) -> bool:
    left_tokens = tokenize_person_name(left)
    right_tokens = tokenize_person_name(right)
    if not left_tokens or not right_tokens:
        return False

    if left_tokens == right_tokens:
        return True

    if not tokens_compatible(left_tokens[0], right_tokens[0]):
        return False
    if not tokens_compatible(left_tokens[-1], right_tokens[-1]):
        return False

    shorter, longer = (
        (left_tokens, right_tokens) if len(left_tokens) <= len(right_tokens) else (right_tokens, left_tokens)
    )
    matched = 0
    for token in shorter:
        if any(tokens_compatible(token, candidate) for candidate in longer):
            matched += 1

    return matched / max(1, len(shorter)) >= 0.75


def score_person_name_completeness(value: str) -> tuple[int, int, int]:
    tokens = tokenize_person_name(value)
    return (
        len(tokens),
        sum(1 for token in tokens if len(token) > 1),
        len(normalize_person_name_for_match(value)),
    )


def choose_best_person_name(*candidates: str) -> str:
    valid_candidates = [candidate.strip() for candidate in candidates if candidate and candidate.strip()]
    if not valid_candidates:
        return ""
    return max(valid_candidates, key=score_person_name_completeness)


def resolve_person_identity(
    *,
    neopi_path: Path | None,
    profiler_path: Path | None,
    anchor_name: str,
) -> tuple[str, list[str]]:
    notes: list[str] = []
    neopi_name = infer_person_name_from_path(neopi_path)
    profiler_name = infer_person_name_from_path(profiler_path)
    anchor_name = clean_text(anchor_name)

    pdf_name = choose_best_person_name(neopi_name, profiler_name)

    if neopi_name and profiler_name and not person_names_match(neopi_name, profiler_name):
        raise ValueError(
            "Os arquivos de NEO PI-R e Perfil comportamental parecem pertencer a pessoas diferentes. "
            f"NEO PI-R: '{neopi_name}' | Perfil: '{profiler_name}'."
        )

    if anchor_name and pdf_name and not person_names_match(anchor_name, pdf_name):
        raise ValueError(
            "Os arquivos enviados parecem pertencer a pessoas diferentes. "
            f"NEO/Perfil: '{pdf_name}' | Âncoras/Diagnóstico: '{anchor_name}'."
        )

    if anchor_name and pdf_name and person_names_match(anchor_name, pdf_name):
        chosen_name = choose_best_person_name(anchor_name, pdf_name)
        if normalize_person_name_for_match(anchor_name) != normalize_person_name_for_match(pdf_name):
            notes.append(
                "Nome conciliado entre PDFs e planilha de Âncoras/Diagnóstico para preservar a forma mais completa."
            )
        return chosen_name, notes

    return pdf_name or anchor_name, notes


def classify_t_score(t_score: int) -> str:
    if t_score <= 34:
        return "muito baixo"
    if t_score <= 44:
        return "baixo"
    if t_score <= 55:
        return "medio"
    if t_score <= 65:
        return "alto"
    return "muito alto"


def is_extreme(category: str) -> bool:
    return category != "medio"


def find_first_file(base_dir: Path, patterns: list[str]) -> Path | None:
    matches: list[Path] = []
    seen: set[str] = set()
    for pattern in patterns:
        for path in sorted(base_dir.glob(pattern)):
            resolved = str(path.resolve())
            if resolved not in seen:
                seen.add(resolved)
                matches.append(path)
    return matches[0] if matches else None


def load_upload_manifest(base_dir: Path) -> dict[str, Path]:
    manifest_path = base_dir / UPLOAD_MANIFEST_FILENAME
    if not manifest_path.exists():
        return {}

    try:
        raw_payload = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return {}

    files_by_category = raw_payload.get("files_by_category") if isinstance(raw_payload, dict) else None
    if not isinstance(files_by_category, dict):
        return {}

    resolved_files: dict[str, Path] = {}
    for category, relative_name in files_by_category.items():
        if not isinstance(category, str) or not isinstance(relative_name, str):
            continue
        candidate = (base_dir / relative_name).resolve()
        if candidate.exists():
            resolved_files[category] = candidate
    return resolved_files


def extract_pdf_pages(path: Path, engine: str = "pypdf") -> list[str]:
    if engine == "pdfplumber":
        with pdfplumber.open(path) as pdf:
            return [clean_text(page.extract_text() or "") for page in pdf.pages]

    reader = PdfReader(str(path))
    return [clean_text(page.extract_text() or "") for page in reader.pages]


def parse_score_from_text(text: str, label: str) -> dict[str, int] | None:
    normalized_text = normalize_text(text)
    normalized_label = normalize_text(label)
    pattern = re.compile(rf"\b{re.escape(normalized_label)}\b\s+(\d+)\s+(\d+)\b")
    match = pattern.search(normalized_text)
    if not match:
        return None
    return {
        "raw_score": int(match.group(1)),
        "t_score": int(match.group(2)),
    }


def find_first_page_index(pages: list[str], required_terms: list[str]) -> int | None:
    normalized_terms = [normalize_text(term) for term in required_terms]
    for index, page_text in enumerate(pages):
        normalized_page = normalize_text(page_text)
        if all(term in normalized_page for term in normalized_terms):
            return index
    return None


def collect_synthesis_text(pages: list[str]) -> str:
    start_index = None
    header_patterns = [
        re.compile(rf"\b{variant}\b")
        for _canonical_name, variants in NEOPI_SYNTHESIS_HEADINGS
        for variant in variants
    ]
    for index, page_text in enumerate(pages):
        normalized_page = normalize_text(page_text)
        if "sintese dos fatores do neo pi-r" in normalized_page and any(
            pattern.search(page_text) for pattern in header_patterns
        ):
            start_index = index
            break

    if start_index is None:
        return ""

    collected_pages: list[str] = []
    for page_text in pages[start_index:]:
        normalized_page = normalize_text(page_text)
        if "sintese dos fatores do neo pi-r" in normalized_page or any(
            pattern.search(page_text) for pattern in header_patterns
        ):
            collected_pages.append(page_text)
            continue
        break

    return " ".join(collected_pages)


def extract_neopi_synthesis(page_text: str) -> dict[str, str]:
    sections: dict[str, str] = {}
    matches: list[tuple[int, int, str]] = []
    for canonical_name, variants in NEOPI_SYNTHESIS_HEADINGS:
        for variant in variants:
            match = re.search(rf"\b{variant}\b", page_text, flags=re.IGNORECASE)
            if match:
                matches.append((match.start(), match.end(), canonical_name))
                break

    matches.sort(key=lambda item: item[0])
    for index, (_start, end_of_heading, canonical_name) in enumerate(matches):
        end = matches[index + 1][0] if index + 1 < len(matches) else len(page_text)
        sections[canonical_name] = clean_text(page_text[end_of_heading:end])
    return sections


def extract_neopi(path: Path) -> dict[str, Any]:
    pages = extract_pdf_pages(path, engine="pypdf")
    score_page_index = find_first_page_index(pages, ["resultados", "escores padronizados t"])
    score_page = pages[score_page_index] if score_page_index is not None else ""
    synthesis_page = collect_synthesis_text(pages)

    domains: list[dict[str, Any]] = []
    facets: list[dict[str, Any]] = []

    for domain_name in NEOPI_DOMAINS:
        scores = parse_score_from_text(score_page, domain_name)
        if not scores:
            continue
        category = classify_t_score(scores["t_score"])
        domains.append(
            {
                "domain": domain_name,
                **scores,
                "category": category,
                "citation_candidate": is_extreme(category),
            }
        )

    for domain_name, facet_names in NEOPI_FACETS.items():
        for facet_name in facet_names:
            scores = parse_score_from_text(score_page, facet_name)
            if not scores:
                continue
            category = classify_t_score(scores["t_score"])
            facets.append(
                {
                    "domain": domain_name,
                    "facet": facet_name,
                    **scores,
                    "category": category,
                    "citation_candidate": is_extreme(category),
                }
            )

    citation_candidates = [
        *[
            {
                "type": "domain",
                "name": domain["domain"],
                "category": domain["category"],
                "t_score": domain["t_score"],
            }
            for domain in domains
            if domain["citation_candidate"]
        ],
        *[
            {
                "type": "facet",
                "name": facet["facet"],
                "domain": facet["domain"],
                "category": facet["category"],
                "t_score": facet["t_score"],
            }
            for facet in facets
            if facet["citation_candidate"]
        ],
    ]

    synthesis_by_domain = extract_neopi_synthesis(synthesis_page)

    return {
        "pages": len(pages),
        "domains": domains,
        "facets": facets,
        "citation_candidates": citation_candidates,
        "synthesis_by_domain": synthesis_by_domain,
        "friendly_synthesis_by_domain": build_friendly_neopi_synthesis_map(synthesis_by_domain),
    }


def extract_prefixed_value(ws: Any, prefix: str) -> str | None:
    normalized_prefix = normalize_text(prefix)
    for row in ws.iter_rows():
        for cell in row:
            value = clean_text(cell.value)
            if not value:
                continue
            normalized_value = normalize_text(value)
            if normalized_value.startswith(normalized_prefix):
                parts = value.split(":", 1)
                if len(parts) == 2:
                    return clean_text(parts[1])
                if value.startswith(prefix):
                    return clean_text(value[len(prefix) :])
                return value
    return None


def find_row_in_column(ws: Any, column: str, marker: str) -> int | None:
    normalized_marker = normalize_text(marker)
    for row in range(1, ws.max_row + 1):
        value = clean_text(ws[f"{column}{row}"].value)
        if value and normalized_marker in normalize_text(value):
            return row
    return None


def collect_column_section(ws: Any, start_marker: str, end_markers: list[str]) -> list[str]:
    start_row = find_row_in_column(ws, "B", start_marker)
    if start_row is None:
        return []

    end_row = ws.max_row + 1
    for marker in end_markers:
        marker_row = find_row_in_column(ws, "B", marker)
        if marker_row and marker_row > start_row:
            end_row = min(end_row, marker_row)

    values: list[str] = []
    for row in range(start_row + 1, end_row):
        value = clean_text(ws[f"B{row}"].value)
        if not value:
            continue
        if normalize_text(value).startswith("pagina "):
            continue
        values.append(value)
    return values


def extract_profiler_scores(ws: Any) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        style = clean_text(ws[f"A{row}"].value)
        score = ws[f"B{row}"].value
        if not style or score is None:
            continue
        percentage = round(float(score) * 100, 2)
        items.append(
            {
                "style": style,
                "score": round(float(score), 4),
                "percentage": percentage,
            }
        )
    return items


def compute_dominant_style_label(scores: list[dict[str, Any]], threshold_pct: float = 6.0) -> str | None:
    """Return the dominant style label, combining multiple styles when scores are close.

    When the second-highest style is within *threshold_pct* percentage points of the
    top score, it is included in the label. Styles are always ordered canonically
    (Executor → Comunicador → Planejador → Analista) regardless of score rank.
    """
    if not scores:
        return None
    style_order = ["Executor", "Comunicador", "Planejador", "Analista"]
    sorted_scores = sorted(scores, key=lambda x: x["percentage"], reverse=True)
    top_pct = sorted_scores[0]["percentage"]
    dominant_names: list[str] = [sorted_scores[0]["style"]]
    for item in sorted_scores[1:]:
        if (top_pct - item["percentage"]) <= threshold_pct:
            dominant_names.append(item["style"])
        else:
            break
    dominant_names.sort(key=lambda s: style_order.index(s) if s in style_order else 99)
    return " ".join(dominant_names)


def canonical_anchor_name(name: str) -> str:
    simplified = clean_text(name).replace("\n", " ")
    simplified = re.sub(r"\s+", " ", simplified).strip()
    aliases = {
        "Administração geral": "Administrativo Geral",
        "Autonomia e Independência": "Autonomia Independência",
        "Segurança e Estabilidade": "Segurança Estabilidade",
        "Técnico Funcional:": "Técnico Funcional",
        "Administração geral:": "Administrativo Geral",
        "Autonomia e Independência:": "Autonomia Independência",
        "Segurança e Estabilidade:": "Segurança Estabilidade",
        "Criatividade Empreendedora:": "Criatividade Empreendedora",
        "Vontade de servir:": "Vontade de Servir",
        "Puro desafio:": "Puro Desafio",
        "Estilo de Vida:": "Estilo de Vida",
        "Autonomia Independência": "Autonomia Independência",
        "Segurança Estabilidade": "Segurança Estabilidade",
        "Criatividade Empreendedora": "Criatividade Empreendedora",
        "Técnico Funcional": "Técnico Funcional",
        "Administrativo Geral": "Administrativo Geral",
        "Vontade de Servir": "Vontade de Servir",
        "Puro Desafio": "Puro Desafio",
        "Estilo de Vida": "Estilo de Vida",
    }
    normalized = normalize_text(simplified)
    for candidate in aliases:
        if normalize_text(candidate) == normalized:
            return aliases[candidate]
    return simplified


def extract_report_anchor_notes(ws: Any) -> list[str]:
    notes: list[str] = []
    mismatches: list[str] = []
    for row in range(3, 11):
        description = clean_text(ws[f"A{row}"].value)
        anchor_name = canonical_anchor_name(clean_text(ws[f"D{row}"].value))
        if not description or not anchor_name:
            continue
        description_anchor = canonical_anchor_name(description.split(":", 1)[0])
        if normalize_text(description_anchor) != normalize_text(anchor_name):
            mismatches.append(f"row {row}: descricao={description_anchor} | rotulo={anchor_name}")
    if mismatches:
        notes.append(
            "Career anchor description rows do not fully align with anchor labels in the "
            "report workbook input sheet: " + "; ".join(mismatches)
        )
    return notes


def extract_report_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True)
    summary_ws = workbook.worksheets[0]
    profiler_ws = workbook.worksheets[2]
    anchor_input_ws = workbook.worksheets[3]
    culture_ws = workbook.worksheets[4]

    sections = {
        "neopi": collect_column_section(summary_ws, "3.1- NEOPI-R", ["3.2- Profiler"]),
        "profiler": collect_column_section(summary_ws, "3.2- Profiler", ["3.3- Âncoras de Carreira"]),
        "career_anchors": collect_column_section(
            summary_ws,
            "3.3- Âncoras de Carreira",
            ["3.4- Diagnóstico Cultural"],
        ),
        "cultural_diagnosis": collect_column_section(
            summary_ws,
            "3.4- Diagnóstico Cultural",
            ["IV. Conclusão"],
        ),
        "conclusion": collect_column_section(summary_ws, "IV. Conclusão", []),
    }

    profiler_scores = extract_profiler_scores(profiler_ws)
    dominant_profiler_label = compute_dominant_style_label(profiler_scores)

    culture_scores: list[dict[str, Any]] = []
    for row in range(3, 7):
        culture_name = clean_text(culture_ws[f"D{row}"].value)
        score = culture_ws[f"E{row}"].value
        if not culture_name or score is None:
            continue
        culture_scores.append(
            {
                "culture": culture_name,
                "score": float(score),
                "description": CULTURE_DESCRIPTIONS.get(culture_name, ""),
            }
        )

    return {
        "person": {
            "name": extract_prefixed_value(summary_ws, "Nome:"),
            "application_date": extract_prefixed_value(summary_ws, "Data da aplicação:"),
            "business_unit": extract_prefixed_value(summary_ws, "Unidade de negócios"),
            "demand": extract_prefixed_value(summary_ws, "Demanda:"),
        },
        "profiler_scores": profiler_scores,
        "dominant_profiler_style": dominant_profiler_label,
        "cultural_scores_from_template": culture_scores,
        "sections": sections,
        "notes": extract_report_anchor_notes(anchor_input_ws),
    }


def extract_anchor_summary(ws: Any) -> list[dict[str, Any]]:
    start_row = None
    for row in range(1, ws.max_row + 1):
        values = [clean_text(cell.value) for cell in ws[row]]
        if any(normalize_text(value) == normalize_text("Âncoras") for value in values if value):
            start_row = row
            break
    if start_row is None:
        return []

    items: list[dict[str, Any]] = []
    for row in range(start_row + 1, start_row + 9):
        name = canonical_anchor_name(clean_text(ws[f"B{row}"].value))
        average = ws[f"D{row}"].value
        if not name or average is None:
            continue
        items.append(
            {
                "name": name,
                "average": float(average),
                "description": CAREER_ANCHOR_DESCRIPTIONS.get(name, ""),
            }
        )
    return items


def extract_culture_summary(ws: Any) -> list[dict[str, Any]]:
    summary_row = None
    for row in range(1, ws.max_row + 1):
        values = [clean_text(cell.value) for cell in ws[row]]
        if any(normalize_text(value) == normalize_text("Minha preferência") for value in values if value):
            summary_row = row
    if summary_row is None:
        return []

    items: list[dict[str, Any]] = []
    for row in range(summary_row + 1, summary_row + 5):
        name = clean_text(ws[f"B{row}"].value)
        score = ws[f"C{row}"].value
        if not name or score is None:
            continue
        canonical_name = canonical_anchor_name(name)
        items.append(
            {
                "culture": canonical_name,
                "score": round(float(score), 4),
                "description": CULTURE_DESCRIPTIONS.get(canonical_name, ""),
            }
        )
    return items


def extract_anchor_and_culture_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True)
    if _looks_like_report_workbook(workbook):
        raise ValueError(
            "A planilha enviada em 'Âncoras e diagnóstico' parece ser o relatório final, e não a planilha de âncoras/diagnóstico."
        )
    if len(workbook.worksheets) < 2:
        raise ValueError(
            "A planilha enviada em 'Âncoras e diagnóstico' não possui a estrutura esperada para Âncoras e Diagnóstico."
        )

    anchors_ws = workbook.worksheets[0]
    culture_ws = workbook.worksheets[1]

    anchor_scores = extract_anchor_summary(anchors_ws)
    top_anchors = sorted(anchor_scores, key=lambda item: item["average"], reverse=True)[:2]

    culture_scores = extract_culture_summary(culture_ws)
    top_cultures = sorted(culture_scores, key=lambda item: item["score"], reverse=True)[:2]

    if not anchor_scores and not culture_scores:
        raise ValueError(
            "Nao foi possivel identificar os dados da planilha de Âncoras e Diagnóstico. Verifique se o arquivo correto foi enviado nesse card."
        )

    return {
        "person": {
            "name": clean_text(anchors_ws["C2"].value),
            "role": clean_text(anchors_ws["C3"].value),
            "application_date": clean_text(anchors_ws["C4"].value),
        },
        "career_anchors": {
            "scores": anchor_scores,
            "top_anchors": top_anchors,
        },
        "cultural_diagnosis": {
            "scores": culture_scores,
            "top_cultures": top_cultures,
        },
    }


def _looks_like_report_workbook(workbook: Any) -> bool:
    sheet_names = [normalize_text(sheet.title) for sheet in workbook.worksheets]
    signature_names = {"sintese", "dados gerais", "input profiler", "input cultura organizacional"}
    return len(signature_names.intersection(sheet_names)) >= 2


def extract_profiler_extended(path: Path) -> dict[str, Any]:
    pages = extract_pdf_pages(path, engine="pdfplumber")
    page_two = pages[1] if len(pages) > 1 else ""
    page_seven = pages[6] if len(pages) > 6 else ""
    page_eight = pages[7] if len(pages) > 7 else ""

    dominant_match = re.search(
        r"Neste momento,.*?está:\s*((?:[A-Za-zÀ-ÿ]+\s*)+?)\s+em\b",
        page_two, flags=re.DOTALL,
    )
    dominant_style = dominant_match.group(1).strip() if dominant_match else None
    profiler_scores = extract_profiler_scores_from_extended(page_two)
    if not dominant_style and profiler_scores:
        dominant_style = compute_dominant_style_label(profiler_scores)

    long_form_sections = {
        "page_2_overview": page_two,
        "page_7_management": page_seven,
        "page_8_sales_and_motivation": page_eight,
    }

    return {
        "pages": len(pages),
        "scores": profiler_scores,
        "dominant_style_from_pdf": dominant_style,
        "long_form_sections": long_form_sections,
    }


def extract_profiler_scores_from_extended(page_text: str) -> list[dict[str, Any]]:
    compact = clean_text(page_text)
    score_match = re.search(
        r"(\d+(?:[.,]\d+)?)%\s+(\d+(?:[.,]\d+)?)%\s+(\d+(?:[.,]\d+)?)%\s+(\d+(?:[.,]\d+)?)%\s+Executor\s+Comunicador\s+Planejador\s+Analista",
        compact,
        flags=re.IGNORECASE,
    )
    if not score_match:
        return []

    style_order = ["Executor", "Comunicador", "Planejador", "Analista"]
    scores: list[dict[str, Any]] = []
    for style_name, raw_percentage in zip(style_order, score_match.groups()):
        percentage = float(raw_percentage.replace(",", "."))
        scores.append(
            {
                "style": style_name,
                "score": round(percentage / 100, 4),
                "percentage": round(percentage, 2),
            }
        )
    return scores


def build_generated_report_sections(
    *,
    neopi_bundle: dict[str, Any],
    profiler_bundle: dict[str, Any],
    anchor_bundle: dict[str, Any],
) -> dict[str, Any]:
    dominant_style = profiler_bundle.get("dominant_style_from_pdf") or "Perfil misto"
    top_anchors = anchor_bundle.get("career_anchors", {}).get("top_anchors", [])
    top_cultures = anchor_bundle.get("cultural_diagnosis", {}).get("top_cultures", [])
    extreme_domains = [item for item in neopi_bundle.get("domains", []) if item.get("citation_candidate")]

    return {
        "neopi": [
            "O Inventario de Personalidade NEO Revisado (NEO PI-R) organiza a leitura do perfil com base nos Cinco Grandes Fatores.",
            "A seguir, estao destacados os fatores com leitura mais relevante para a analise atual.",
            *_build_generated_neopi_lines(neopi_bundle),
        ],
        "profiler": [
            "Ferramenta para mapeamento de estilo comportamental baseada em quatro caracteristicas: Executor, Comunicador, Planejador e Analista.",
            f"Nesse momento, apresenta o estilo: {dominant_style}",
            build_generated_profiler_paragraph(dominant_style),
        ],
        "career_anchors": [
            "O questionario de ancoras de carreira ajuda a identificar valores e motivadores profissionais mais presentes neste momento.",
            *[
                f"{item['name']}: {item.get('description', '')}".strip()
                for item in top_anchors
                if item.get("name")
            ],
        ],
        "cultural_diagnosis": [
            "O diagnostico cultural aponta os ambientes organizacionais com maior aderencia percebida pela pessoa avaliada.",
            *[
                f"{item['culture']}: {item.get('description', '')}".strip()
                for item in top_cultures
                if item.get("culture")
            ],
        ],
        "conclusion": build_generated_conclusion_lines(
            dominant_style=dominant_style,
            top_anchors=top_anchors,
            top_cultures=top_cultures,
            extreme_domains=extreme_domains,
        ),
    }


def _build_generated_neopi_lines(neopi_bundle: dict[str, Any]) -> list[str]:
    synthesis_map = neopi_bundle.get("friendly_synthesis_by_domain") or neopi_bundle.get("synthesis_by_domain", {})
    lines: list[str] = []
    for domain_name in NEOPI_DOMAINS:
        synthesis = clean_text(synthesis_map.get(domain_name))
        if synthesis:
            lines.append(f"{domain_name}: {rewrite_neopi_synthesis_text(synthesis, domain_name=domain_name)}")
            continue

        matching_domain = next(
            (item for item in neopi_bundle.get("domains", []) if item.get("domain") == domain_name),
            None,
        )
        if matching_domain:
            lines.append(
                f"{domain_name}: resultado classificado como {matching_domain.get('category')} com T score {matching_domain.get('t_score')}."
            )
    return lines


def build_generated_profiler_paragraph(dominant_style: str) -> str:
    summaries = {
        "Executor": "Tende a atuar com energia para acao, senso de urgencia e disposicao para assumir desafios.",
        "Comunicador": "Tende a se comunicar com fluidez, fortalecer relacoes e mobilizar pessoas com entusiasmo.",
        "Planejador": "Costuma atuar com prudencia, consistencia e preferencia por previsibilidade na execucao.",
        "Analista": "Tende a priorizar profundidade, criterio e qualidade tecnica nas entregas.",
    }
    # For multi-style labels (e.g. "Comunicador Planejador"), combine each style's
    # individual description so all named styles are represented in the paragraph.
    individual_styles = [s.strip() for s in dominant_style.split() if s.strip() in summaries]
    if individual_styles:
        return " ".join(summaries[s] for s in individual_styles)
    return summaries.get(
        dominant_style,
        "O estilo predominante contribui para a leitura do comportamento no contexto profissional.",
    )


def build_generated_conclusion_lines(
    *,
    dominant_style: str,
    top_anchors: list[dict[str, Any]],
    top_cultures: list[dict[str, Any]],
    extreme_domains: list[dict[str, Any]],
) -> list[str]:
    lines: list[str] = []
    if dominant_style:
        lines.append(f"Predominio do estilo {dominant_style} no contexto avaliado.")
    if top_anchors:
        lines.append("Ancoras mais presentes: " + ", ".join(item["name"] for item in top_anchors if item.get("name")) + ".")
    if top_cultures:
        lines.append(
            "Maior aderencia cultural a "
            + ", ".join(item["culture"] for item in top_cultures if item.get("culture"))
            + "."
        )
    if extreme_domains:
        lines.append(
            "No NEO PI-R, destacam-se "
            + ", ".join(
                _format_domain_category_label(str(item["domain"]), str(item["category"]))
                for item in extreme_domains
                if item.get("domain") and item.get("category")
            )
            + "."
        )
    raw_lines = lines[:5]
    # Prefix each line with a number so build_template_fallback can detect them
    # as conclusion items (it checks for patterns like "1-", "2-" etc.).
    return [f"{i + 1}- {line}" for i, line in enumerate(raw_lines)]


def _format_domain_category_label(domain: str, category: str) -> str:
    feminine_domains = {"Extroversão", "Abertura", "Amabilidade", "Conscienciosidade"}
    adjusted_category = category
    if domain in feminine_domains:
        adjusted_category = (
            category.replace("muito baixo", "muito baixa")
            .replace("baixo", "baixa")
            .replace("muito alto", "muito alta")
            .replace("alto", "alta")
        )
    return f"{domain} {adjusted_category}"


def build_bundle(base_dir: Path) -> dict[str, Any]:
    manifest_files = load_upload_manifest(base_dir)
    files = {
        key: manifest_files.get(key) or find_first_file(base_dir, patterns)
        for key, patterns in REPORT_FILE_PATTERNS.items()
    }
    missing = [key for key in REQUIRED_FILE_KEYS if files.get(key) is None]
    if missing:
        raise FileNotFoundError(f"Missing expected files for keys: {', '.join(missing)}")

    neopi_bundle = extract_neopi(files["neopi_pdf"])
    profiler_bundle = extract_profiler_extended(files["profiler_pdf"])
    report_workbook = extract_report_workbook(files["report_workbook"]) if files.get("report_workbook") else None
    anchor_bundle = extract_anchor_and_culture_workbook(files["anchors_workbook"])
    generated_sections = build_generated_report_sections(
        neopi_bundle=neopi_bundle,
        profiler_bundle=profiler_bundle,
        anchor_bundle=anchor_bundle,
    )
    resolved_person_name, identity_notes = resolve_person_identity(
        neopi_path=files.get("neopi_pdf"),
        profiler_path=files.get("profiler_pdf"),
        anchor_name=anchor_bundle["person"].get("name") or "",
    )

    notes = []
    if report_workbook:
        notes.extend(report_workbook.get("notes", []))
    else:
        notes.append("Report workbook not provided. Generated sections were created from the raw assessment files.")
    notes.extend(identity_notes)
    if (
        (report_workbook or {}).get("dominant_profiler_style") or profiler_bundle.get("dominant_style_from_pdf")
    ) and anchor_bundle["career_anchors"]["top_anchors"]:
        notes.append(
            "The sample suggests a concise final report can be built from a small subset "
            "of signals: dominant profiler style, top-2 anchors, top-2 cultures, and NEO "
            "non-medium indicators."
        )

    person = {
        "name": (report_workbook or {}).get("person", {}).get("name") or resolved_person_name,
        "application_date": (report_workbook or {}).get("person", {}).get("application_date")
        or anchor_bundle["person"].get("application_date"),
        "business_unit": (report_workbook or {}).get("person", {}).get("business_unit"),
        "demand": (report_workbook or {}).get("person", {}).get("demand"),
        "role": anchor_bundle["person"].get("role"),
    }

    return {
        "input_dir": str(base_dir.resolve()),
        "files": {key: str(va