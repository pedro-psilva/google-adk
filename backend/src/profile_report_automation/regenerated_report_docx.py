from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.section import _Header
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from profile_report_automation.regenerated_report_pdf import (
    ANCHORS_INTRO_TEXT,
    CULTURE_INTRO_TEXT,
    METHODOLOGY_TEXT,
    NEOPI_INTRO_TEXTS,
    PROFILER_INTRO_TEXT,
    _build_report_context,
    _resolve_logo_path,
    _wrap_anchor_label,
)


ACCENT_ORANGE = RGBColor(0xFF, 0x5A, 0x1F)
TEXT_COLOR = RGBColor(0x11, 0x11, 0x11)
LOGO_WIDTH_CM = 1.72
CHART_WIDTH_CM = 15.4
TITLE_SIZE_PT = 19
HEADING_SIZE_PT = 13
SUBHEADING_SIZE_PT = 11.5
BODY_SIZE_PT = 10.5
LOGO_PADDING_PX = 24

CHART_TEXT = "#111111"
CHART_BORDER = "#b5b5b5"
CHART_GRID = "#d0d0d0"
ANCHOR_BAR_COLOR = "#7db3c8"
CULTURE_LINE_COLOR = "#1b75bc"
SCALE = 3

FONT_CANDIDATES = {
    "regular": [
        "C:/Windows/Fonts/arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ],
    "bold": [
        "C:/Windows/Fonts/arialbd.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf",
    ],
}


def build_regenerated_docx_report(
    path: str | Path,
    bundle: dict[str, Any],
    draft: dict[str, Any],
    *,
    logo_path: str | Path | None = None,
) -> None:
    target_path = Path(path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_document(document, logo_path=logo_path)

    context = _build_report_context(bundle, draft)

    _add_title(document, "ANÁLISE DE PERFIL COMPORTAMENTAL")
    _add_section_heading(document, "I. Identificação:")
    for label, value in context["person_rows"]:
        _add_label_value_paragraph(document, label, value)

    _add_spacer(document, 4)
    _add_section_heading(document, "II. Metodologia e objetivo:")
    _add_body_paragraph(document, METHODOLOGY_TEXT)

    _add_section_heading(document, "III. Resultados da Análise:")
    _add_subheading(document, "3.1- NEOPI-R")
    for paragraph in NEOPI_INTRO_TEXTS:
        _add_body_paragraph(document, paragraph)
    for line in context["neopi_lines"]:
        _add_labeled_text_paragraph(document, line)

    document.add_page_break()

    _add_section_heading(document, "III. Resultados da Análise:")
    _add_subheading(document, "3.2- Profiler")
    _add_body_paragraph(document, PROFILER_INTRO_TEXT)
    _add_chart_picture(document, _build_profiler_chart_image(context["profiler_scores"]), width_cm=CHART_WIDTH_CM)
    _add_label_value_paragraph(document, "Nesse momento, apresenta o estilo", f"{context['profiler_dominant_style']}.")
    _add_body_paragraph(document, context["profiler_summary"])

    _add_subheading(document, "3.3- Âncoras de Carreira")
    _add_body_paragraph(document, ANCHORS_INTRO_TEXT)
    _add_chart_picture(document, _build_anchor_chart_image(context["anchor_scores"]), width_cm=CHART_WIDTH_CM)
    for line in context["anchor_lines"]:
        _add_labeled_text_paragraph(document, line)

    document.add_page_break()

    _add_section_heading(document, "III. Resultados da Análise:")
    _add_subheading(document, "3.4- Diagnóstico Cultural")
    _add_body_paragraph(document, CULTURE_INTRO_TEXT)
    _add_chart_picture(document, _build_culture_chart_image(context["culture_values"]), width_cm=11.6)
    for line in context["culture_lines"]:
        _add_labeled_text_paragraph(document, line)

    _add_section_heading(document, "IV. Conclusão:")
    _add_body_paragraph(document, "A partir dos indicadores de seu perfil, apresenta:")
    for index, line in enumerate(context["conclusion_lines"], start=1):
        _add_body_paragraph(document, f"{index}- {line}")

    document.save(target_path)


def _configure_document(document: Document, *, logo_path: str | Path | None = None) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.28)
    section.footer_distance = Cm(0.45)
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(BODY_SIZE_PT)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.line_spacing = 1.0

    resolved_logo_path = _resolve_logo_path(logo_path)
    if resolved_logo_path is None:
        return

    logo_bytes = _build_logo_stream(resolved_logo_path).getvalue()
    _configure_header_with_logo(section.header, logo_bytes)
    _configure_header_with_logo(section.first_page_header, logo_bytes)
    _configure_header_with_logo(section.even_page_header, logo_bytes)


def _add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(9)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(TITLE_SIZE_PT)
    run.font.color.rgb = TEXT_COLOR


def _add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    square_run = paragraph.add_run("■ ")
    square_run.font.name = "Arial"
    square_run.font.size = Pt(HEADING_SIZE_PT)
    square_run.font.color.rgb = ACCENT_ORANGE

    text_run = paragraph.add_run(text)
    text_run.bold = True
    text_run.font.name = "Arial"
    text_run.font.size = Pt(HEADING_SIZE_PT)
    text_run.font.color.rgb = TEXT_COLOR


def _add_subheading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.underline = True
    run.font.name = "Arial"
    run.font.size = Pt(SUBHEADING_SIZE_PT)
    run.font.color.rgb = TEXT_COLOR


def _add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2.5)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(BODY_SIZE_PT)
    run.font.color.rgb = TEXT_COLOR


def _add_label_value_paragraph(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)

    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.name = "Arial"
    label_run.font.size = Pt(BODY_SIZE_PT)
    label_run.font.color.rgb = TEXT_COLOR

    value_run = paragraph.add_run(value)
    value_run.font.name = "Arial"
    value_run.font.size = Pt(BODY_SIZE_PT)
    value_run.font.color.rgb = TEXT_COLOR


def _add_labeled_text_paragraph(document: Document, raw_text: str) -> None:
    label, separator, remainder = raw_text.partition(":")
    if not separator:
        _add_body_paragraph(document, raw_text)
        return

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)

    label_run = paragraph.add_run(f"{label}:")
    label_run.bold = True
    label_run.font.name = "Arial"
    label_run.font.size = Pt(BODY_SIZE_PT)
    label_run.font.color.rgb = TEXT_COLOR

    remainder_run = paragraph.add_run(remainder)
    remainder_run.font.name = "Arial"
    remainder_run.font.size = Pt(BODY_SIZE_PT)
    remainder_run.font.color.rgb = TEXT_COLOR


def _add_chart_picture(document: Document, chart_stream: BytesIO, *, width_cm: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(4)

    chart_stream.seek(0)
    run = paragraph.add_run()
    run.add_picture(chart_stream, width=Cm(width_cm))


def _add_spacer(document: Document, height_pt: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(height_pt)


def _build_logo_stream(path: Path) -> BytesIO:
    original = Image.open(path).convert("RGBA")
    alpha_channel = original.getchannel("A")
    bbox = alpha_channel.getbbox()

    if bbox is None:
        flattened = Image.new("RGB", original.size, "white")
        flattened.paste(original, mask=alpha_channel)
        return _image_to_stream(flattened)

    left, top, right, bottom = bbox
    left = max(0, left - LOGO_PADDING_PX)
    top = max(0, top - LOGO_PADDING_PX)
    right = min(original.width, right + LOGO_PADDING_PX)
    bottom = min(original.height, bottom + LOGO_PADDING_PX)

    cropped = original.crop((left, top, right, bottom))
    flattened = Image.new("RGB", cropped.size, "white")
    flattened.paste(cropped, mask=cropped.getchannel("A"))
    return _image_to_stream(flattened)


def _configure_header_with_logo(header: _Header, logo_bytes: bytes) -> None:
    _clear_header_paragraphs(header)
    paragraph = header.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    run.add_picture(BytesIO(logo_bytes), width=Cm(LOGO_WIDTH_CM))


def _clear_header_paragraphs(header: _Header) -> None:
    for paragraph in list(header.paragraphs):
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def _build_profiler_chart_image(scores: list[dict[str, Any]]) -> BytesIO:
    width = 392 * SCALE
    height = 112 * SCALE
    image, draw = _new_chart_canvas(width, height)

    label_font = _load_font(8 * SCALE)
    small_font = _load_font(7 * SCALE)
    value_font = _load_font(8 * SCALE, bold=True)

    legend_y = 24 * SCALE
    legend_x = 96 * SCALE
    for item in scores:
        color = _hex_color(item.get("color")) or "#666666"
        style = str(item.get("style") or "")
        draw.rectangle((legend_x, legend_y, legend_x + 7 * SCALE, legend_y + 7 * SCALE), fill=color, outline=CHART_TEXT, width=2)
        _draw_text(draw, legend_x + 12 * SCALE, legend_y - 1 * SCALE, style, label_font, CHART_TEXT)
        legend_x += 56 * SCALE if style != "Comunicador" else 68 * SCALE

    chart_x = 42 * SCALE
    chart_y = 44 * SCALE
    chart_width = 304 * SCALE
    chart_height = 24 * SCALE
    axis_max = 120.0
    total_width = chart_width * (100.0 / axis_max)
    current_x = chart_x

    for item in scores:
        percentage = float(item.get("percentage") or 0.0)
        segment_width = total_width * (percentage / 100.0)
        segment_box = (
            int(round(current_x)),
            chart_y,
            int(round(current_x + segment_width)),
            chart_y + chart_height,
        )
        draw.rectangle(segment_box, fill=_hex_color(item.get("color")) or "#666666", outline=CHART_TEXT, width=2)
        if segment_width >= 22 * SCALE:
            _draw_centered_text(
                draw,
                current_x + (segment_width / 2),
                chart_y + (chart_height / 2),
                _format_percentage(percentage),
                value_font,
                CHART_TEXT,
            )
        current_x += segment_width

    draw.line((chart_x, chart_y + chart_height, chart_x + chart_width, chart_y + chart_height), fill=CHART_TEXT, width=2)
    draw.line((chart_x, chart_y, chart_x, chart_y + chart_height), fill=CHART_TEXT, width=2)

    for tick in range(0, 121, 20):
        tick_x = chart_x + (chart_width * (tick / axis_max))
        draw.line((tick_x, chart_y + chart_height, tick_x, chart_y + chart_height + 5 * SCALE), fill=CHART_TEXT, width=2)
        _draw_centered_text(draw, tick_x, chart_y + chart_height + 18 * SCALE, f"{tick},00%", small_font, CHART_TEXT)

    return _image_to_stream(image)


def _build_anchor_chart_image(scores: list[dict[str, Any]]) -> BytesIO:
    width = 392 * SCALE
    height = 188 * SCALE
    image, draw = _new_chart_canvas(width, height)

    label_font = _load_font(7 * SCALE)
    axis_font = _load_font(8 * SCALE)
    value_font = _load_font(8 * SCALE)

    chart_left = 140 * SCALE
    chart_bottom = 46 * SCALE
    chart_width = 170 * SCALE
    chart_height = 104 * SCALE
    chart_top = chart_bottom + chart_height
    max_value = 6.0
    row_height = 12 * SCALE
    bar_height = 8 * SCALE

    for tick in range(0, 7):
        tick_x = chart_left + ((chart_width / max_value) * tick)
        draw.line((tick_x, chart_bottom, tick_x, chart_top), fill=CHART_GRID, width=2)
        _draw_centered_text(draw, tick_x, chart_top + 14 * SCALE, str(tick), axis_font, CHART_TEXT)

    draw.line((chart_left, chart_top, chart_left + chart_width, chart_top), fill=CHART_TEXT, width=2)

    for index, item in enumerate(scores):
        value = float(item.get("average") or 0.0)
        row_y = chart_bottom + (index * row_height)
        label_lines = _wrap_anchor_label(str(item.get("name") or "")).splitlines()
        label_height = len(label_lines) * 8 * SCALE
        label_y = row_y + (bar_height / 2) - (label_height / 2)
        for offset, line in enumerate(label_lines):
            _draw_right_text(draw, chart_left - 12 * SCALE, label_y + (offset * 8 * SCALE), line, label_font, CHART_TEXT)

        bar_width = chart_width * (value / max_value)
        draw.rectangle(
            (
                chart_left,
                row_y,
                int(round(chart_left + bar_width)),
                row_y + bar_height,
            ),
            fill=ANCHOR_BAR_COLOR,
            outline=CHART_TEXT,
            width=2,
        )
        _draw_text(draw, chart_left + bar_width + 6 * SCALE, row_y - 1 * SCALE, _format_decimal(value), value_font, CHART_TEXT)

    return _image_to_stream(image)


def _build_culture_chart_image(values: list[float]) -> BytesIO:
    width = 276 * SCALE
    height = 222 * SCALE
    image, draw = _new_chart_canvas(width, height)

    label_font = _load_font(9 * SCALE)
    axis_font = _load_font(8 * SCALE)

    center_x = 132 * SCALE
    center_y = 104 * SCALE
    radius = 76 * SCALE
    max_value = 40.0
    levels = [10, 20, 30, 40]

    for level in levels:
        level_radius = radius * (level / max_value)
        polygon = [
            (center_x, center_y - level_radius),
            (center_x + level_radius, center_y),
            (center_x, center_y + level_radius),
            (center_x - level_radius, center_y),
        ]
        draw.polygon(polygon, outline=CHART_GRID, width=2)

    draw.line((center_x, center_y - radius, center_x, center_y + radius), fill=CHART_GRID, width=2)
    draw.line((center_x - radius, center_y, center_x + radius, center_y), fill=CHART_GRID, width=2)

    for level in [0, *levels]:
        level_radius = radius * (level / max_value)
        _draw_text(draw, center_x - 22 * SCALE, center_y - level_radius - 3 * SCALE, str(level), axis_font, CHART_TEXT)

    _draw_centered_text(draw, center_x, center_y - radius - 16 * SCALE, "Clã", label_font, CHART_TEXT)
    _draw_text(draw, center_x + radius + 10 * SCALE, center_y - 4 * SCALE, "Inovativa", label_font, CHART_TEXT)
    _draw_centered_text(draw, center_x, center_y + radius + 14 * SCALE, "Mercado", label_font, CHART_TEXT)
    _draw_right_text(draw, center_x - radius - 10 * SCALE, center_y - 4 * SCALE, "Hierárquica", label_font, CHART_TEXT)

    polygon_points = []
    vectors = [(0, -1), (1, 0), (0, 1), (-1, 0)]
    for value, (x_vector, y_vector) in zip(values, vectors):
        scaled_radius = radius * (min(max(float(value), 0.0), max_value) / max_value)
        polygon_points.append((center_x + (scaled_radius * x_vector), center_y + (scaled_radius * y_vector)))

    draw.line([*polygon_points, polygon_points[0]], fill=CULTURE_LINE_COLOR, width=5, joint="curve")
    for point_x, point_y in polygon_points:
        draw.ellipse((point_x - 4 * SCALE, point_y - 4 * SCALE, point_x + 4 * SCALE, point_y + 4 * SCALE), fill=CULTURE_LINE_COLOR)

    return _image_to_stream(image)


def _new_chart_canvas(width: int, height: int) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((0, 0, width - 1, height - 1), radius=10 * SCALE, outline=CHART_BORDER, width=2, fill="white")
    return image, draw


def _load_font(size: int, *, bold: bool = False) -> ImageFont.ImageFont | ImageFont.FreeTypeFont:
    key = "bold" if bold else "regular"
    for candidate in FONT_CANDIDATES[key]:
        font_path = Path(candidate)
        if font_path.exists():
            return ImageFont.truetype(str(font_path), size=size)
    return ImageFont.load_default()


def _draw_text(
    draw: ImageDraw.ImageDraw,
    x: float,
    y: float,
    text: str,
    font: ImageFont.ImageFont | ImageFont.FreeTypeFont,
    fill: str,
) -> None:
    draw.text((int(round(x)), int(round(y))), text, font=font, fill=fill)


def _draw_centered_text(
    draw: ImageDraw.ImageDraw,
    center_x: float,
    center_y: float,
    text: str,
    font: ImageFont.ImageFont | ImageFont.FreeTypeFont,
    fill: str,
) -> None:
    left, top, right, bottom = draw.textbbox((0, 0), text, font=font)
    width = right - left
    height = bottom - top
    draw.text((center_x - (width / 2), center_y - (height / 2)), text, font=font, fill=fill)


def _draw_right_text(
    draw: ImageDraw.ImageDraw,
    right_x: float,
    y: float,
    text: str,
    font: ImageFont.ImageFont | ImageFont.FreeTypeFont,
    fill: str,
) -> None:
    left, top, right, bottom = draw.textbbox((0, 0), text, font=font)
    width = right - left
    draw.text((right_x - width, y), text, font=font, fill=fill)


def _image_to_stream(image: Image.Image) -> BytesIO:
    stream = BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _hex_color(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    hex_value = getattr(value, "hexval", None)
    if callable(hex_value):
        raw_hex = str(hex_value())
        if raw_hex.startswith("0x"):
            return f"#{raw_hex[2:]}"
        return raw_hex
    rgb = getattr(value, "rgb", None)
    if rgb is None:
        return None
    red, green, blue = [int(round(channel * 255)) for channel in rgb()]
    return f"#{red:02x}{green:02x}{blue:02x}"


def _format_percentage(value: float) -> str:
    return f"{value:.2f}%".replace(".", ",")


def _format_decimal(value: float) -> str:
    text = f"{value:.1f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")
